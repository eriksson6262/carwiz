#!/usr/bin/env python3
"""Nouzový režim (19.10.2026): Renteon nepřihlášen -> CZ.json/SK.json se sestaví
z posledního importu/exportu (PL), pravidel dle stavu v projektovém dokumentu (RULES) a
R odvozeného z DC Carwiz přes poslední kalibraci (R = (DC - b)/k). Obsazenost neznámá (u=None)."""
import json, os, sys, datetime as dt
from openpyxl import load_workbook
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(__file__))
from reprice import CLASSES, BANDS, LOC

PL_SRC = {'CZ': os.environ['PL_CZ'], 'SK': os.environ['PL_SK']}
# kalibrace 12.10.2026 (procedure doc)
CAL = {('PRG', 'E'): (1.33, 24), ('PRG', 'C'): (1.28, 22), ('PRG', 'I'): (1.27, 20), ('PRG', 'F'): (1.24, 9),
       ('OSR', 'C'): (1.09, 9), ('OSR', 'I'): (1.11, 2), ('OSR', 'F'): (1.12, 2), ('OSR', 'E'): (1.09, 9),
       ('BTS', 'C'): (1.18, 12), ('BTS', 'I'): (1.19, 8), ('BTS', 'E'): (1.18, 12), ('BTS', 'F'): (1.19, 8)}
VAT = {'CZ': 1.21, 'SK': 1.23}


def read_pl(path):
    ws = load_workbook(path)['Ceny']; R = list(ws.iter_rows(values_only=True)); h = list(R[0])
    out = []
    for r in R[1:]:
        if not r[0]: continue
        d = dict(zip(h, r)); off = d['Office']
        out.append(dict(S=d['ServiceName'], O=off, Oid=26 if off == 'OSR' else None,
                        F=d['DateFrom'].strftime('%d.%m.%Y.'), T=d['DateTo'].strftime('%d.%m.%Y.'),
                        D={b: d[b] for b in BANDS}, disc=d.get('DiscountPercentage')))
    return out


def rules(country, sipps):
    def R(i, adj, uf, ut, cats=None, force=False, pf=None, pt=None):
        return dict(Id=i, PriceAdjustmentPercentage=adj, UtilizationPercentageFrom=uf, UtilizationPercentageTo=ut,
                    IsAllCarCategories=cats is None, CarCategoriesText=', '.join(cats) if cats else '', Durations='1-',
                    ForceApply=force, PeriodFrom=pf, PeriodTo=pt, ValidTo=None, Modification='stav 12.10.2026 (Renteon nepřihlášen – nepřečteno)')
    if country == 'CZ':
        c = [s for s in sipps if s not in ('GFAR', 'IFAR')]
        return [R(None, -30, 0, 50, c), R(None, -15, 51, 70, c), R(None, 15, 86, 94, c), R(None, 50, 95, 100, c),
                R(None, 100, 51, 100, ['GFAR']), R(3760, 500, None, None, None, True, '24.12.2026.', '27.12.2026.')]
    return [R(None, -60, 0, 70), R(None, 100, 85, 100), R(2307, 300, None, None, None, True, '24.12.2026.', '27.12.2026.')]


dc = json.load(open(f'{BASE}/data/DC.json'))
SIPP2K = {s: k for k, l in CLASSES.items() for s in l}
for country in ('CZ', 'SK'):
    PL = read_pl(PL_SRC[country]); sipps = sorted({r['S'] for r in PL})
    RA = []
    for r in dc['res']:
        L = r['L']; c, office, _ = LOC[L]
        if c != country: continue
        m = {}
        for s, v in r['cw'].items():
            k = SIPP2K.get(s)
            if not k: continue
            kk, b = CAL.get((L, k), (1.3, 0.0))
            R = (v - b) / kk
            if R > 0: m[s] = [round(R, 2), None]
        RA.append(dict(o=office, p=r['p'], du=r['du'], m=m))
    json.dump(dict(PL=PL, RA=RA, RULES=rules(country, sipps), FA={'25': {}, '26': {}, '128': {}}, MODE='fallback'),
              open(f'{BASE}/data/{country}.json', 'w'))
    print(country, len(PL), len(RA))
