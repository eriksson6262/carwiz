#!/usr/bin/env python3
"""Kontrola anomálií před importem (Erik 20.10.2026). Čte IMPORT xlsx (list Ceny), živý ceník (data/CZ.json, SK.json),
obsazenost (FA), vzorky konkurence (data/DC|RC|CT|WZ.json) a předchozí běh (env PREV_DIR: dc_*.json, ct_*.json, wizz_*.json, rc_*.json).
Výstup out/anomalies.json: {CZ:{block:[...], warn:[...]}, SK:{...}, common:{warn:[...]}}.
Použití: IMPORT_CZ=… IMPORT_SK=… RUN_TODAY=YYYY-MM-DD python3 tools/check_anomalies.py"""
import json, os, sys, math, glob, statistics as st, datetime as dt
from collections import defaultdict
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(__file__))
from reprice import BASE, D, ISO, fmt, TODAY, SIPP2K, BANDS, BMIN, LOC, expand

BAND_LEN = ['2-3', '4-6', '7-7', '8-13', '14-14', '15-20']
DATA = {c: json.load(open(f'{BASE}/data/{c}.json')) for c in ('CZ', 'SK')}
MIN_AFTER_C = {'CZ': {'E': 5.0, 'C': 5.5, 'I': 6.5, 'F': 13.0}, 'SK': {'E': 6.5, 'C': 7.5, 'I': 9.0, 'F': 13.0}}
MIN_SIPP = {'GFAR': 65.0}
MIN_1DAY = 35.0
OFF = {'CZ': {None: '25', 26: '26'}, 'SK': {None: '128'}}
LNAME = {'25': 'Praha', '26': 'Ostrava', '128': 'Bratislava'}
CHN = {'DC': 'DiscoverCars', 'RC': 'Rentalcars.com', 'CT': 'CarTrawler easyJet', 'WZ': 'CarTrawler Wizz Air'}


def max_disc(c, sipp):
    m = 1.0
    for r in DATA[c]['RULES']:
        adj = r.get('PriceAdjustmentPercentage') or 0
        if adj >= 0 or r.get('ForceApply'): continue
        cats = [x.strip() for x in (r.get('CarCategoriesText') or '').split(',')]
        if r.get('IsAllCarCategories') or not r.get('CarCategoriesText') or sipp in cats:
            m = min(m, 1 + adj / 100)
    return m


def read_import(path):
    ws = load_workbook(path)['Ceny']; R = list(ws.iter_rows(values_only=True)); h = list(R[0]); out = []
    for r in R[1:]:
        if not r or r[0] is None: continue
        d = dict(zip(h, r)); off = d.get('Office')
        out.append(dict(S=d['ServiceName'], Oid=26 if off == 'OSR' else None, F=d['DateFrom'], T=d['DateTo'], D={b: d.get(b) for b in BANDS}))
    for r in out:
        for k in ('F', 'T'):
            v = r[k]; r[k] = v.date() if isinstance(v, dt.datetime) else (D(v) if isinstance(v, str) else v)
    return out


def fleet(c):
    """{(Oid, SIPP): (max vozů, rezervací v 60 dnech, [(den, booked, total) přebookováno])}"""
    res = {}
    for oid, key in OFF[c].items():
        for sipp, v in DATA[c]['FA'].get(key, {}).items():
            if not v['T'] or sipp.startswith(('Celk',)): continue
            d0 = D(v['d0'])
            over = [(d0 + dt.timedelta(days=i), b, t) for i, (t, b) in enumerate(zip(v['T'], v['B'])) if (b or 0) > (t or 0)]
            res[(oid, sipp)] = (max(v['T']), sum((b or 0) for b in v['B'][:60]), over)
    return res


def live_at(PL, sipp, oid, d):
    for r in PL:
        if r['S'] == sipp and r['Oid'] == oid and D(r['F']) <= d <= D(r['T']):
            return r['D']


def check_country(c, path, channels):
    B, W = [], []
    PL = DATA[c]['PL']
    if not PL: B.append('chybí živý ceník z Renteonu')
    if not DATA[c].get('FA'): B.append('chybí obsazenost (Fleet availability) z Renteonu')
    locs = [L for L, v in LOC.items() if v[0] == c]
    nch = sum(1 for ch, d in channels.items() if any(r['L'] in locs and any(r['comp'].values()) for r in d['res']))
    if nch < 2: B.append(f'k dispozici jen {nch} kanál(y) konkurence (min. 2)')
    imp = read_import(path)
    FL = fleet(c)
    # 1) formální kontrola: čísla, pásma, třídy, spojitost
    bad = [(r['S'], r['F']) for r in imp if any(not isinstance(r['D'].get(b), (int, float)) or r['D'].get(b) is None or r['D'][b] <= 0 for b in BANDS)]
    if bad: B.append(f'nečíselné/prázdné ceny v {len(bad)} řádcích, např. {bad[:3]}')
    live_keys = {(r['S'], r['Oid']) for r in PL if D(r['T']) >= TODAY}
    imp_keys = {(r['S'], r['Oid']) for r in imp}
    miss = sorted(live_keys - imp_keys, key=str)
    if miss: B.append(f'v importu chybí třída/pobočka oproti živému ceníku: {miss}')
    live_bands = {b for r in PL for b in r['D']}
    if live_bands - set(BANDS): B.append(f'živý ceník má pásma, která import nemá: {sorted(live_bands - set(BANDS))}')
    grp = defaultdict(list)
    for r in imp: grp[(r['S'], r['Oid'])].append((r['F'], r['T']))
    for key, lst in grp.items():
        lst.sort()
        if lst[0][0] != TODAY: B.append(f'{key}: první řádek začíná {lst[0][0]}, ne {TODAY}')
        for (a1, b1), (a2, b2) in zip(lst, lst[1:]):
            if a2 != b1 + dt.timedelta(days=1): B.append(f'{key}: mezera/překryv {b1} → {a2}')
        lend = max(D(r['T']) for r in PL if (r['S'], r['Oid']) == key) if key in live_keys else None
        if lend and lst[-1][1] != lend: B.append(f'{key}: konec {lst[-1][1]} ≠ živý ceník {lend}')
    # 2) minima, €/den s délkou, Vánoce
    low, low1, rise_rep, rise_oth, uniform, xmas = [], [], [], [], [], []
    for r in imp:
        d = r['D']
        if any(not isinstance(d.get(b), (int, float)) for b in BANDS): continue
        L = LNAME[OFF[c][r['Oid']]]
        mn = MIN_SIPP.get(r['S']) or MIN_AFTER_C[c].get('E' if r['S'][0] == 'M' else r['S'][0], 6.5)   # mini = cena EDMR (Erik 20.10.) -> minimum jako E
        md = max_disc(c, r['S'])
        for b in BANDS[:7]:
            if d[b] * md < mn - 0.011: low.append(f"{L} {r['S']} {r['F']:%d.%m.}–{r['T']:%d.%m.} {b} {d[b]} € (po slevě {d[b]*md:.2f} < {mn})")
        if d['1-1'] * md < MIN_1DAY - 0.011: low1.append(f"{L} {r['S']} {r['F']:%d.%m.}–{r['T']:%d.%m.} 1-1 {d['1-1']} € (po slevě {d['1-1']*md:.2f} < 35)")
        rises = [f'{b1}<{b2}' for b1, b2 in zip(BAND_LEN, BAND_LEN[1:]) if d[b2] > d[b1] + 0.011]
        o0 = live_at(PL, r['S'], r['Oid'], r['F'])
        changed = o0 is None or any(abs((o0.get(b) or 0) - d[b]) > 0.011 for b in BANDS)
        if rises:
            (rise_rep if (r['S'] in SIPP2K and changed) else rise_oth).append(f"{L} {r['S']} {r['F']:%d.%m.%Y}–{r['T']:%d.%m.%Y} ({', '.join(rises)})")
        if (r['Oid'], r['S']) in FL and FL[(r['Oid'], r['S'])][0] > 0 and len({round(d[b], 2) for b in BAND_LEN}) == 1:
            uniform.append(f"{L} {r['S']} {r['F']:%d.%m.}–{r['T']:%d.%m.} ({d['7-7']} €)")
        if r['F'] >= dt.date(2026, 12, 24) and r['T'] <= dt.date(2026, 12, 27):
            o = live_at(PL, r['S'], r['Oid'], r['F'])
            if o is None or any(abs((o.get(b) or 0) - d[b]) > 0.011 for b in BANDS):
                xmas.append(f"{L} {r['S']} {r['F']:%d.%m.}–{r['T']:%d.%m.}")
    if low: B.append(f'cena pod minimem po slevě ({len(low)}×): ' + '; '.join(low[:5]))
    if low1: B.append(f'1 den pod 35 € po slevě ({len(low1)}×): ' + '; '.join(low1[:5]))
    if rise_rep: B.append(f'€/den roste s délkou u přeceňovaných tříd ({len(rise_rep)}×): ' + '; '.join(rise_rep[:5]))
    if xmas: B.append(f'změněné řádky 24.–27.12. ({len(xmas)}×): ' + '; '.join(xmas[:5]))
    if rise_oth: W.append(f'€/den roste s délkou u nepřeceňovaných/nezměněných řádků ({len(rise_oth)}×): ' + '; '.join(rise_oth[:6]))
    if uniform: W.append(f'jednotná cena pro všechny délky ({len(uniform)}×): ' + '; '.join(uniform[:6]))
    # 3) velké změny proti živému ceníku
    big30 = []; nrows = 0; nchg30 = 0; maxchg = {}
    for r in imp:
        o = live_at(PL, r['S'], r['Oid'], r['F'])
        if not o: continue
        nrows += 1
        rel = [abs(r['D'][b] / o[b] - 1) for b in BAND_LEN if o.get(b) and isinstance(r['D'].get(b), (int, float))]
        if rel and max(rel) > 0.30: nchg30 += 1
        has_fleet = FL.get((r['Oid'], r['S']), (0,))[0] > 0
        if not has_fleet: continue
        d = r['F']
        while d <= min(r['T'], TODAY + dt.timedelta(days=30)):
            o2 = live_at(PL, r['S'], r['Oid'], d)
            if o2:
                for b in ['1-1'] + BAND_LEN:
                    if o2.get(b) and isinstance(r['D'].get(b), (int, float)):
                        ch = r['D'][b] / o2[b] - 1
                        if abs(ch) > 0.40:
                            k = (r['S'], r['Oid'], b)
                            if k not in maxchg or abs(ch) > abs(maxchg[k][0]): maxchg[k] = (ch, d, o2[b], r['D'][b])
            d += dt.timedelta(days=1)
    for (s, oid, b), (ch, d, o, n) in sorted(maxchg.items(), key=lambda x: str(x[0])):
        big30.append(f"{LNAME[OFF[c][oid]]} {s} {b} {d:%d.%m.} {o}→{n} € ({100*ch:+.0f} %)")
    if big30: W.append(f'změna > ±40 % u třídy s flotilou v příštích 30 dnech – neblokuje, odpovídá cíli podle konkurence (Erik 20.10.) ({len(big30)}×): ' + '; '.join(big30[:8]))
    if nrows and nchg30 / nrows > 0.40: B.append(  # Erik 20.10.2026 (17:00): blokuje až > 40 % řádků
        f'{nchg30}/{nrows} řádků ({100*nchg30/nrows:.0f} %) změněno o víc než 30 % – podezřelá data')
    # 4) obsazenost
    for (oid, s), (n, b60, over) in sorted(FL.items(), key=lambda x: str(x[0])):
        L = LNAME[OFF[c][oid]]
        if n > 0 and b60 == 0: W.append(f'{L} {s}: {n} vozů bez jediné rezervace v příštích 60 dnech')
        if over:
            ov = [x for x in over if x[2] > 0]; z = [x for x in over if not x[2]]
            if ov: W.append(f"{L} {s}: přebookováno {ov[0][0]:%d.%m.}–{ov[-1][0]:%d.%m.} (až {max(x[1] for x in ov)}/{max(x[2] for x in ov)})")
            if z: W.append(f"{L} {s}: rezervace bez flotily {z[0][0]:%d.%m.}–{z[-1][0]:%d.%m.} (až {max(x[1] for x in z)})")
    nofleet = sorted({f"{LNAME[OFF[c][r['Oid']]]} {r['S']}" for r in imp if FL.get((r['Oid'], r['S']), (0,))[0] == 0})
    if nofleet: W.append('třídy bez flotily s cenou (výplňové ceny, nepřeceňují se): ' + ', '.join(nofleet))
    return dict(block=B, warn=W)


def comp_change(cur, prev, name):
    """konkurent výrazně zlevnil/zdražil: medián poměru nejlevnější konkurent nyní/minule za (pobočka, skupina, délka), shodné termíny."""
    out = []
    P = {(r['L'], r['p'], r['du'], k): v[0][0] for r in prev['res'] for k, v in r['comp'].items() if v}
    by = defaultdict(list)
    for r in cur['res']:
        for k, v in r['comp'].items():
            q = P.get((r['L'], r['p'], r['du'], k))
            if v and q: by[(r['L'], k)].append(v[0][0] / q)
    for (L, k), v in sorted(by.items()):
        if len(v) >= 5:
            m = st.median(v)
            if abs(m - 1) > 0.20: out.append(f'{name} {L} {k}: nejlevnější konkurent {100*(m-1):+.0f} % proti minulému běhu (medián {len(v)} termínů)')
    return out


def supplier_change(cur, prev, name):
    out = []
    def sup(x): return {(r['L'], s[1]) for r in x['res'] for v in r['comp'].values() for s in v}
    new = sup(cur) - sup(prev)
    if new: out.append(f'{name}: noví konkurenti {sorted(new)}')
    return out


if __name__ == '__main__':
    channels = {}
    for ch in ('DC', 'RC', 'CT', 'WZ'):
        p = f'{BASE}/data/{ch}.json'
        if os.path.exists(p): channels[ch] = expand(json.load(open(p)))
    res = {c: check_country(c, os.environ[f'IMPORT_{c}'], channels) for c in ('CZ', 'SK')}
    common = []
    # Carwiz chybí na kanálu
    for ch, d in channels.items():
        for L in ('PRG', 'OSR', 'BTS'):
            rr = [r for r in d['res'] if r['L'] == L]
            if not rr: common.append(f'{CHN[ch]} {L}: žádná data'); continue
            n = sum(1 for r in rr if r['cw'])
            if n / len(rr) < 0.5: common.append(f'{CHN[ch]} {L}: Carwiz zobrazen jen v {n}/{len(rr)} vzorcích')
    # nové třídy / pobočky proti minulému běhu Renteonu
    prevdir = os.environ.get('PREV_DIR')
    if prevdir:
        for c in ('CZ', 'SK'):
            pp = sorted(glob.glob(f'{prevdir}/renteon_{c}_*.json'))
            if pp:
                pr = json.load(open(pp[-1]))
                a = {(r['S'], r['Oid']) for r in pr['PL']}; b = {(r['S'], r['Oid']) for r in DATA[c]['PL']}
                if b - a: common.append(f'{c}: nové třídy/pobočky v ceníku {sorted(b - a, key=str)}')
                fa = set(pr.get('FA', {}).keys()); fb = set(DATA[c]['FA'].keys())
                if fb - fa: common.append(f'{c}: nové pobočky v obsazenosti {sorted(fb - fa)}')
        for ch, pref in (('DC', 'dc'), ('RC', 'rc'), ('CT', 'ct'), ('WZ', 'wizz')):
            pp = sorted(glob.glob(f'{prevdir}/{pref}_*.json'))
            if ch in channels and pp:
                prev = expand(json.load(open(pp[-1])))
                common += comp_change(channels[ch], prev, CHN[ch]) + supplier_change(channels[ch], prev, CHN[ch])
    res['common'] = dict(warn=common)
    json.dump(res, open(f'{BASE}/out/anomalies.json', 'w'), ensure_ascii=False, indent=1, default=str)
    for c in ('CZ', 'SK'):
        print(c, 'BLOKUJE' if res[c]['block'] else 'OK')
        for x in res[c]['block']: print('  ✖', x)
        for x in res[c]['warn']: print('  !', x)
    for x in common: print('  !', x)
