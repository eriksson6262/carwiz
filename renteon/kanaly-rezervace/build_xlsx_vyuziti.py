#!/usr/bin/env python3
"""Excel „Využití flotily“ – data/aktivity_CZ.csv, data/vyfakturuj_dlouhodobe.csv, output/mesice.csv
→ output/Vyuziti_flotily.xlsx (procenta a součty jako vzorce). Bez osobních údajů."""
import pathlib
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

B = pathlib.Path(__file__).parent
OUT = B / 'output' / 'Vyuziti_flotily.xlsx'
F = 'Arial'
H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3A5F')
INP = Font(name=F, color='0000FF'); TXT = Font(name=F); BOLD = Font(name=F, bold=True)
LINK = Font(name=F, color='008000'); TITLE = Font(name=F, bold=True, size=14)
NOTE = Font(name=F, italic=True, color='666666', size=9)
THIN = Border(bottom=Side(style='thin', color='CCCCCC'))
PCT = '0.0%;-0.0%;-'; NUM = '#,##0;-#,##0;-'; DEC = '#,##0.0;-#,##0.0;-'

CIN = [('kratkodobe', 'Krátkodobé'), ('dlouhodobe', 'Dlouhodobé'), ('interni', 'Interní'), ('asistence', 'Asistence'),
       ('hoppygo', 'HoppyGo'), ('servis_presun', 'Servis / přesun'), ('volne', 'Volné')]


def header(ws, row, labels, widths=None):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row, i, t); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def main():
    akt = pd.read_csv(B / 'data' / 'aktivity_CZ.csv'); akt = akt[akt.flotila == 'pronajem']
    piv = akt.pivot_table(index=['pob', 'mesic'], columns='cinnost', values='dny', aggfunc='sum', fill_value=0).reset_index()
    vf = pd.read_csv(B / 'data' / 'vyfakturuj_dlouhodobe.csv')
    mes = pd.read_csv(B / 'output' / 'mesice.csv')

    wb = Workbook()

    # ---------- Využití (dny) ----------
    ws = wb.active; ws.title = 'Využití'
    ws['A1'] = 'Využití flotily pronájmu podle činností aut (dny aut)'; ws['A1'].font = TITLE
    ws['A2'] = ('Zdroj: Renteon CZ, činnosti aut (api/carActivity/getSimpleListForCar) rozpočítané po dnech, flotila „Pronájem vozového parku“; '
                'deaktivovaná (prodaná/vrácená) auta se nepočítají. Modře = data, černě = vzorce.'); ws['A2'].font = NOTE
    cols = ['Pobočka', 'Měsíc'] + [n for _, n in CIN] + ['Celkem dní', 'Ø aut'] + [n + ' %' for _, n in CIN] + ['Obsazeno %']
    header(ws, 4, cols, [9, 10] + [11] * 7 + [11, 8] + [11] * 7 + [11])
    r0 = 5
    for i, row in enumerate(piv.sort_values(['pob', 'mesic']).itertuples(index=False)):
        r = r0 + i; d = row._asdict()
        ws.cell(r, 1, d['pob']).font = INP; ws.cell(r, 2, d['mesic']).font = INP
        for j, (k, _) in enumerate(CIN):
            c = ws.cell(r, 3 + j, round(float(d.get(k, 0)), 1)); c.font = INP; c.number_format = DEC
        ws.cell(r, 10, f'=SUM(C{r}:I{r})').number_format = DEC
        ws.cell(r, 11, f'=IFERROR(J{r}/DAY(EOMONTH(DATEVALUE(B{r}&"-01"),0)),0)').number_format = DEC
        for j in range(7):
            col = get_column_letter(3 + j)
            ws.cell(r, 12 + j, f'=IFERROR({col}{r}/$J{r},0)').number_format = PCT
        ws.cell(r, 19, f'=IFERROR((C{r}+D{r}+E{r}+F{r}+G{r})/J{r},0)').number_format = PCT
        for c in range(1, 20):
            ws.cell(r, c).border = THIN
            if ws.cell(r, c).font != INP:
                ws.cell(r, c).font = TXT
    last = r0 + len(piv) - 1
    ws.cell(last + 2, 1, 'Obsazeno = krátkodobé + dlouhodobé + interní + asistence + HoppyGo. Servis/přesun a volné se nepočítají. '
            'Volné (Líný) zahrnuje i obrátku mezi pronájmy a auta držená pro rezervace.').font = NOTE
    ws.freeze_panes = 'C5'
    U = (r0, last)

    # ---------- Obsazenost Renteon ----------
    wo = wb.create_sheet('Obsazenost Renteon')
    wo['A1'] = 'Obsazenost podle Renteonu (Kontrola vozu → dostupnost flotily)'; wo['A1'].font = TITLE
    wo['A2'] = ('Zdroj: Renteon CZ api/carControl/fleetAvailability, denní hodnoty zprůměrované za měsíc. „Obsazeno“ = BookedFleetCount / TotalFleet '
                '(počítá i rezervace a blokace), „Krátkodobé“ = BookedByRentCount / Total. Pro historii chybí auta, která už byla deaktivována.'); wo['A2'].font = NOTE
    header(wo, 4, ['Pobočka', 'Měsíc', 'Obsazeno % (Renteon)', 'Z toho krátkodobé %', 'Ostatní obsazení %'], [9, 10, 16, 16, 16])
    ob = {'PRG': [(83, 10), (83, 12), (86, 21), (82, 16), (84, 16), (84, 20), (80, 17), (77, 25), (83, 39), (87, 46), (89, 51), (82, 42)],
          'OSR': [(80, 46), (74, 38), (70, 33), (68, 13), (73, 24), (60, 12), (76, 25), (83, 22), (72, 22), (92, 48), (90, 52), (78, 29)]}
    months = ['2025-10', '2025-11', '2025-12', '2026-01', '2026-02', '2026-03', '2026-04', '2026-05', '2026-06', '2026-07', '2026-08', '2026-09']
    r = 5
    for p, vals in ob.items():
        for m, (o, k) in zip(months, vals):
            # OSR: obsazeno = BookedCount/Total pobočky (BookedFleet je za celou firmu)
            wo.cell(r, 1, p).font = INP; wo.cell(r, 2, m).font = INP
            c = wo.cell(r, 3, o / 100); c.font = INP; c.number_format = PCT
            c = wo.cell(r, 4, k / 100); c.font = INP; c.number_format = PCT
            c = wo.cell(r, 5, f'=C{r}-D{r}'); c.font = TXT; c.number_format = PCT
            for cc in range(1, 6): wo.cell(r, cc).border = THIN
            r += 1
    wo.cell(5, 3).comment = Comment('PRG: průměr BookedFleetCount/TotalFleet za dny měsíce.', 'Claude')
    wo.cell(17, 3).comment = Comment('OSR: průměr BookedCount/Total za pobočku (BookedFleet v datech Renteonu odpovídá celé firmě).', 'Claude')
    r += 1
    wo.cell(r, 1, 'Praha VII–IX 2026, denní špička').font = BOLD; r += 1
    PEAK = r
    for lab, v, f in [('Nejvyšší obsazenost dne (bez anomálie 16. 7. = 103 %)', 0.94, PCT), ('Dní s obsazeností ≥ 90 %', 29, NUM), ('Dní s obsazeností ≥ 95 %', 1, NUM)]:
        wo.cell(r, 1, lab).font = TXT; c = wo.cell(r, 3, v); c.font = INP; c.number_format = f; r += 1
    wo.freeze_panes = 'C5'

    # ---------- Vyfakturuj ----------
    wv = wb.create_sheet('Vyfakturuj dlouhodobé')
    wv['A1'] = 'Dlouhodobé pronájmy fakturované přes Vyfakturuj (Kč bez DPH)'; wv['A1'].font = TITLE
    wv['A2'] = ('Zdroj: Vyfakturuj RX Auto s.r.o., export dokladů podle DUZP (faktury, daňové doklady, opravné, příjmové), stornované vynechány; '
                'jen nájemní položky (bez pokut, jistin, prodejů aut). Jednotlivé firmy nejsou uvedeny.'); wv['A2'].font = NOTE
    header(wv, 4, ['Měsíc', 'Firemní dlouhodobí (bez HAVEX, Fábián)', 'Fábián Automotive', 'HAVEX Mobility', 'Global Assistance (OSR)', 'Celkem', 'Aut na dokladech (SPZ)'],
           [10, 20, 16, 16, 18, 14, 14])
    r = 5
    for row in vf.itertuples(index=False):
        wv.cell(r, 1, row.mesic).font = INP
        for j, v in enumerate([row.ostatni_dlouhodobi, row.fabian, row.havex, row.asistence_osr]):
            c = wv.cell(r, 2 + j, int(v)); c.font = INP; c.number_format = NUM
        c = wv.cell(r, 6, f'=SUM(B{r}:E{r})'); c.font = TXT; c.number_format = NUM
        c = wv.cell(r, 7, int(row.aut_spz)); c.font = INP; c.number_format = NUM
        for cc in range(1, 8): wv.cell(r, cc).border = THIN
        r += 1
    vl = r - 1
    wv.cell(r, 1, 'Celkem').font = BOLD
    for j in range(2, 7):
        col = get_column_letter(j); c = wv.cell(r, j, f'=SUM({col}5:{col}{vl})'); c.font = BOLD; c.number_format = NUM
    wv.cell(r + 1, 1, 'Ø měsíc s fakturací').font = BOLD
    for j in range(2, 7):
        col = get_column_letter(j); c = wv.cell(r + 1, j, f'=IFERROR(AVERAGEIF({col}5:{col}{vl},">0"),0)'); c.font = BOLD; c.number_format = NUM
    wv.cell(5, 4).comment = Comment('HAVEX Mobility: klasifikace položek nejistá (zápočty, přefakturace) – brát orientačně.', 'Claude')
    wv.cell(r + 3, 1, 'Fábián od IV/2026. Aut na dokladech = různé SPZ v nájemních položkách daného měsíce (včetně aut mimo Renteon).').font = NOTE
    VTOT = r

    # ---------- Výnos na den ----------
    wd = wb.create_sheet('Výnos na den')
    wd['A1'] = 'Krátkodobý vs dlouhodobý pronájem – výnos na den auta (Praha)'; wd['A1'].font = TITLE
    wd['A2'] = 'Krátkodobé: výnos smluv z Renteonu (bez škod a pokut) / dny krátkodobých činností. Dlouhodobé: typická měsíční sazba firemních smluv ve Vyfakturuj.'; wd['A2'].font = NOTE
    header(wd, 4, ['Měsíc', 'Výnos krátkodobé (Kč)', 'Dny krátkodobé', 'Kč / den'], [10, 18, 14, 12])
    r = 5
    prg = mes[mes.pob_out == 'PRG'].sort_values('mesic')
    for row in prg.itertuples(index=False):
        wd.cell(r, 1, row.mesic).font = INP
        c = wd.cell(r, 2, int(row.vynos)); c.font = INP; c.number_format = NUM
        c = wd.cell(r, 3, f"=SUMIFS('Využití'!$C${U[0]}:$C${U[1]},'Využití'!$A${U[0]}:$A${U[1]},\"PRG\",'Využití'!$B${U[0]}:$B${U[1]},A{r})")
        c.font = LINK; c.number_format = DEC
        c = wd.cell(r, 4, f'=IFERROR(B{r}/C{r},0)'); c.font = TXT; c.number_format = NUM
        for cc in range(1, 5): wd.cell(r, cc).border = THIN
        r += 1
    dl = r - 1
    wd.cell(r, 1, '12 měsíců').font = BOLD
    for j, f in [(2, f'=SUM(B5:B{dl})'), (3, f'=SUM(C5:C{dl})'), (4, f'=IFERROR(B{r}/C{r},0)')]:
        c = wd.cell(r, j, f); c.font = BOLD; c.number_format = NUM
    ST = r
    wd.cell(5, 2).comment = Comment('Výnos podle měsíce výdeje auta, dny podle skutečných dnů v měsíci – po měsících se mírně posouvá, celkem za rok sedí.', 'Claude')
    r += 2
    wd.cell(r, 1, 'Dlouhodobé – předpoklad').font = BOLD; r += 1
    wd.cell(r, 1, 'Sazba / auto / měsíc – dolní (Kč)').font = TXT; c = wd.cell(r, 2, 12000); c.font = INP; c.number_format = NUM; c.fill = PatternFill('solid', fgColor='FFFF00'); lo = r; r += 1
    wd.cell(r, 1, 'Sazba / auto / měsíc – horní (Kč)').font = TXT; c = wd.cell(r, 2, 17000); c.font = INP; c.number_format = NUM; c.fill = PatternFill('solid', fgColor='FFFF00'); hi = r; r += 1
    wd.cell(lo, 3, 'Typické sazby firemních smluv ve Vyfakturuj X/2025–IX/2026 (bez HAVEX, Fábián).').font = NOTE
    wd.cell(r, 1, 'Kč / den – dolní').font = TXT; c = wd.cell(r, 2, f'=B{lo}*12/365'); c.number_format = NUM; c.font = TXT; r += 1
    wd.cell(r, 1, 'Kč / den – horní').font = TXT; c = wd.cell(r, 2, f'=B{hi}*12/365'); c.number_format = NUM; c.font = TXT; r += 1
    wd.cell(r, 1, 'Krátkodobé vs dlouhodobé (horní)').font = BOLD; c = wd.cell(r, 2, f'=IFERROR(D{ST}/B{r-1},0)'); c.number_format = '0.0"×"'; c.font = BOLD
    for col, w in zip('ABCD', [34, 18, 14, 12]): wd.column_dimensions[col].width = w

    # ---------- Souhrn ----------
    s = wb.create_sheet('Souhrn', 0)
    s['A1'] = 'Využití flotily – souhrn (X/2025–IX/2026)'; s['A1'].font = TITLE
    s['A2'] = 'RX Auto s.r.o., Renteon CZ + Vyfakturuj. Všechna čísla níže jsou vzorce z ostatních listů.'; s['A2'].font = NOTE
    header(s, 4, ['Praha – flotila pronájmu', 'Srpen 2026', 'Leden–březen 2026 (Ø)'], [34, 16, 22])
    rng = lambda col: f"'Využití'!${col}${U[0]}:${col}${U[1]}"
    def share(col, months):
        num = '+'.join(f'SUMIFS({rng(col)},{rng("A")},"PRG",{rng("B")},"{m}")' for m in months)
        den = '+'.join(f'SUMIFS({rng("J")},{rng("A")},"PRG",{rng("B")},"{m}")' for m in months)
        return f'=IFERROR(({num})/({den}),0)'
    r = 5
    def avg_cars(months):
        parts = '+'.join(f'SUMIFS({rng("K")},{rng("A")},"PRG",{rng("B")},"{m}")' for m in months)
        return f'=({parts})/{len(months)}'
    s.cell(r, 1, 'Ø aut').font = TXT
    for j, ms in [(2, ['2026-08']), (3, ['2026-01', '2026-02', '2026-03'])]:
        c = s.cell(r, j, avg_cars(ms)); c.font = LINK; c.number_format = NUM
    r += 1
    for col, (k, nm) in zip('CDEFGHI', CIN):
        if k in ('asistence', 'hoppygo'):
            continue
        s.cell(r, 1, nm).font = TXT
        for j, ms in [(2, ['2026-08']), (3, ['2026-01', '2026-02', '2026-03'])]:
            c = s.cell(r, j, share(col, ms)); c.font = LINK; c.number_format = PCT
        for cc in range(1, 4): s.cell(r, cc).border = THIN
        r += 1
    r += 1
    header(s, r, ['Obsazenost podle Renteonu (Praha)', 'Hodnota', '']); r += 1
    so = "'Obsazenost Renteon'"
    for lab, f, fmt in [('Ø červenec–srpen 2026', f"=AVERAGE({so}!C14:C15)", PCT),
                        ('Nejvyšší obsazenost dne', f"={so}!C{PEAK}", PCT), ('Dní ≥ 90 % (VII–IX)', f"={so}!C{PEAK + 1}", NUM)]:
        s.cell(r, 1, lab).font = TXT; c = s.cell(r, 2, f); c.font = LINK; c.number_format = fmt; r += 1
    r += 1
    header(s, r, ['Vyfakturuj – dlouhodobé (Kč bez DPH)', 'Ø měsíc', '12 měsíců']); r += 1
    sv = "'Vyfakturuj dlouhodobé'"
    for lab, col in [('Firemní dlouhodobí (bez HAVEX, Fábián)', 'B'), ('Fábián Automotive (od IV/2026)', 'C')]:
        s.cell(r, 1, lab).font = TXT
        c = s.cell(r, 2, f'={sv}!{col}{VTOT + 1}'); c.font = LINK; c.number_format = NUM
        c = s.cell(r, 3, f'={sv}!{col}{VTOT}'); c.font = LINK; c.number_format = NUM
        r += 1
    r += 1
    header(s, r, ['Výnos na den auta (Praha)', 'Kč / den', '']); r += 1
    sd = "'Výnos na den'"
    s.cell(r, 1, 'Krátkodobý pronájem (12 měsíců)').font = TXT; c = s.cell(r, 2, f'={sd}!D{ST}'); c.font = LINK; c.number_format = NUM; r += 1
    s.cell(r, 1, 'Dlouhodobý – horní sazba').font = TXT; c = s.cell(r, 2, f'={sd}!B{hi + 2}'); c.font = LINK; c.number_format = NUM; r += 1
    r += 1
    s.cell(r, 1, 'Závěry').font = BOLD; r += 1
    for t in ['V zimě drží flotilu dlouhodobé pronájmy, v létě nesou tržby krátkodobé – zvážit uvolnění části dlouhodobých aut na léto nebo sezónní sazby.',
              'Interní („Vnitřní“) je trvale ~20 aut flotily pronájmu + ~6 HG – ověřit, co to je; pokud nevydělávají, je to největší rezerva.',
              'Volné v létě je hlavně obrátka mezi pronájmy a auta držená pro rezervace, ne stojící auta.',
              'Renteon „obsazeno“ zahrnuje i budoucí rezervace a blokace, proto je vyšší než podíl skutečně pronajatých dní.']:
        c = s.cell(r, 1, '• ' + t); c.font = TXT; c.alignment = Alignment(wrap_text=True, vertical='top'); s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        s.row_dimensions[r].height = 30; r += 1

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
