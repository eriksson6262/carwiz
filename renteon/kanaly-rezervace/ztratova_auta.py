#!/usr/bin/env python3
"""Seznam ztrátových aut 2025–2026: data/pl_auta.csv + data/dny_auta.csv → output/Ztratova_auta.xlsx.
GFAR: náklad podle Erika 40 tis. Kč/měs vč. pojištění + 30 tis. Kč/rok servis (v datech chybí leasing)."""
import pathlib
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

B = pathlib.Path(__file__).parent
F = 'Arial'; H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3A5F')
INP = Font(name=F, color='0000FF'); TXT = Font(name=F); BOLD = Font(name=F, bold=True); TITLE = Font(name=F, bold=True, size=14)
NOTE = Font(name=F, italic=True, color='666666', size=9); RED = PatternFill('solid', fgColor='FDE2E1'); YEL = PatternFill('solid', fgColor='FFFF00')
THIN = Border(bottom=Side(style='thin', color='DDDDDD')); NUM = '#,##0;[Red]-#,##0;-'; PCT = '0%;-0%;-'

pl = pd.read_csv(B / 'data' / 'pl_auta.csv'); dn = pd.read_csv(B / 'data' / 'dny_auta.csv')
d = pl[pl.rok.isin([2025, 2026])].merge(dn, on=['spz', 'rok'], how='left')
d['sipp'] = d.sipp.fillna(d.sipp_renteon)
for c in ['dny_kratkodobe', 'dny_dlouhodobe', 'dny_interni', 'dny_volne']:
    d[c] = d[c].fillna(0)
d['dny'] = d.dny_kratkodobe + d.dny_dlouhodobe + d.dny_interni + d.dny_volne
d['gfar'] = d.sipp == 'GFAR'
d['leasing_vse'] = d.leasing  # pl_auta.leasing už obsahuje Carnetto půjčovné, pojisteni i Carnetto ČGP
KURZ = 24.3
sk = pd.read_csv(B / 'data' / 'vyfakturuj_sk_spz.csv'); sk = sk[sk.rok.isin([2025, 2026])]
d = d.merge(sk.groupby(['spz', 'rok']).vynos_eur.sum().rename('v_sk_eur').reset_index(), on=['spz', 'rok'], how='left')
d['v_sk_eur'] = d.v_sk_eur.fillna(0)
cz = pd.read_csv(B / 'data' / 'vyfakturuj_cz_spz.csv'); cz = cz[cz.rok.isin([2025, 2026]) & (cz.vynos_czk > 0)]
prodeje = pd.read_csv(B / 'data' / 'prodeje_spz.csv').set_index('spz')
fmt = lambda v: f'{v:,.0f}'.replace(',', ' ')
def poznamka(spz, dl_dny):
    out = []
    for src, df, val, cur in [('Vyfakturuj SK', sk, 'vynos_eur', '€'), ('Vyfakturuj CZ', cz, 'vynos_czk', 'Kč')]:
        x = df[df.spz == spz]
        if len(x):
            odb = '; '.join(sorted({o for v in x.odberatel for o in str(v).split('; ')}))
            out.append(f"Dlouhodobě přes {src}: {odb}; {int(x.mesicu.sum())} měs. fakturováno, {fmt(x[val].sum())} {cur} bez DPH")
    if spz in prodeje.index:
        r = prodeje.loc[spz]; out.append(f"Prodáno – doklad {r.odberatel} ({fmt(r.doklad_kc)} Kč), prodej není ve výsledku")
    if not out and dl_dny > 30:
        out.append(f"V Renteonu {dl_dny:.0f} dní dlouhodobě, ve Vyfakturuj CZ ani SK žádná nájemní faktura se SPZ – ověřit")
    return ' | '.join(out)

ren = d[d.sipp.notna() & ~d.spz.str.startswith('(') & (d.dny >= 30)]
ren = ren[(ren.poznamka.fillna('') != 'bez leasingu v datech') | ren.gfar]
a = ren.groupby('spz').agg(model=('model', 'first'), sipp=('sipp', 'first'), pobocka=('pobocka', 'last'), dny=('dny', 'sum'),
                           kr=('dny_kratkodobe', 'sum'), dl=('dny_dlouhodobe', 'sum'), iv=('dny_interni', 'sum'), vo=('dny_volne', 'sum'), v_kr=('vynos_renteon', 'sum'),
                           v_dl=('vynos_vyfakturuj', 'sum'), v_hx=('vynos_havex', 'sum'), v_sk=('v_sk_eur', 'sum'), leas=('leasing_vse', 'sum'), poj=('pojisteni', 'sum'),
                           serv=('servis_opravy', 'sum'), vys=('vysledek', 'sum'), gfar=('gfar', 'max')).reset_index()
# GFAR: náklady podle zadání
g = a.gfar
a.loc[g, 'leas'] = (a.loc[g, 'dny'] * 40000 * 12 / 365).round(0); a.loc[g, 'poj'] = 0; a.loc[g, 'serv'] = (a.loc[g, 'dny'] * 30000 / 365).round(0)
a['v_dl'] = a.v_dl + (a.v_sk * KURZ).round(0)   # + Vyfakturuj RX Auto SK (v P&L po autech chyběl)
a['vys'] = a.v_kr + a.v_dl + a.v_hx - a.leas - a.poj - a.serv
a['pozn'] = [poznamka(x.spz, x.dl) for x in a.itertuples()]
a['prodano'] = a.spz.isin(prodeje.index)
a = a[a.vys < 0].copy()
a['vys_mes'] = a.vys / a.dny * 30.4
a = a.sort_values('vys')

wb = Workbook(); ws = wb.active; ws.title = 'Ztrátová auta'
ws['A1'] = 'Ztrátová auta 2025–2026 (I–X)'; ws['A1'].font = TITLE
ws['A2'] = ('Kč bez DPH. Výnos: Renteon (krátkodobé) + Vyfakturuj RX Auto CZ a SK (SK v € × 24,3) a HAVEX (dlouhodobé). Náklady: leasing HAVEX/EFS + Carnetto přefakturace, pojištění, servis/opravy '
            '(dashboard „Náklady vs výnosy“). Bez mezd, provizí brokerů, franšízy a paliva. Jen auta v Renteonu s ≥ 30 dny provozu a s leasingem v datech. '
            'Seřazeno od největší ztráty. Modře = data, černě = vzorce.'); ws['A2'].font = NOTE; ws['A2'].alignment = Alignment(wrap_text=True); ws.merge_cells('A2:T2'); ws.row_dimensions[2].height = 42
cols = ['SPZ', 'Model', 'SIPP', 'Pobočka', 'Dní v provozu', 'Krátkodobě %', 'Dlouhodobě %', 'Interní + volné %', 'Výnos krátkodobé', 'Výnos dlouhodobé (Vyfakturuj CZ+SK)', 'Výnos HAVEX',
        'Výnos celkem', 'Leasing', 'Pojištění', 'Servis / opravy', 'Náklady celkem', 'Výsledek', 'Výsledek / měsíc', 'Doporučení', 'Poznámka – dlouhodobé / prodej']
widths = [10, 26, 7, 8, 9, 10, 10, 10, 12, 12, 11, 12, 11, 10, 10, 12, 12, 11, 34, 70]
for i, (t, w) in enumerate(zip(cols, widths), 1):
    c = ws.cell(4, i, t); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[4].height = 32
r = 5
for x in a.itertuples(index=False):
    vals = [x.spz, x.model, x.sipp, x.pobocka, round(x.dny)]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = INP
    bts = x.pobocka == 'BTS'
    for j, v in [(6, x.kr), (7, x.dl), (8, x.iv + x.vo)]:
        c = ws.cell(r, j, None if bts else f'=IFERROR({round(v, 1)}/E{r},0)'); c.number_format = PCT
    for j, v in zip(range(9, 12), [x.v_kr, x.v_dl, x.v_hx]):
        c = ws.cell(r, j, int(v)); c.font = INP; c.number_format = NUM
    ws.cell(r, 12, f'=SUM(I{r}:K{r})').number_format = NUM
    for j, v in zip(range(13, 16), [x.leas, x.poj, x.serv]):
        c = ws.cell(r, j, int(v)); c.font = INP; c.number_format = NUM
        if x.gfar: c.fill = YEL
    ws.cell(r, 16, f'=SUM(M{r}:O{r})').number_format = NUM
    ws.cell(r, 17, f'=L{r}-P{r}').number_format = NUM
    ws.cell(r, 18, f'=IFERROR(Q{r}/E{r}*30.4,0)').number_format = NUM
    kr, dl, ivo = x.kr / x.dny, x.dl / x.dny, (x.iv + x.vo) / x.dny
    rec = ('Prodáno – posoudit včetně prodejní ceny' if x.prodano else
           'Krátce ve flotile – náklady a výnosy nemusí časově sedět' if x.dny < 90 else
           'Bratislava – dny z SK Renteonu nejsou v datech, posoudit podle výnosu' if bts and x.vys_mes > -8000 else
           'Dlouhodobý pronájem bez výnosu v datech – ověřit, jestli se fakturuje' if dl > 0.3 and x.v_dl + x.v_hx < 0.2 * (x.leas + x.poj + x.serv) else
           'Vrátit / prodat' if x.vys_mes < -8000 else
           'Přecenit dlouhodobou smlouvu nebo vrátit do krátkodobého' if dl > 0.5 else
           'Většinu času interní / volné – zjistit proč' if ivo > 0.6 else
           'Málo krátkodobých dní – přesunout na silnější pobočku / do sezóny' if kr < 0.4 else
           'Zkontrolovat náklady (servis, škody)')
    ws.cell(r, 19, rec)
    c = ws.cell(r, 20, x.pozn); c.alignment = Alignment(wrap_text=True, vertical='top')
    if x.gfar:
        ws.cell(r, 13).comment = Comment('GFAR: náklad podle Erika – 40 tis. Kč/měs vč. pojištění + 30 tis. Kč/rok servis, rozpočítáno na dny v provozu.', 'Claude')
    for j in range(1, 21):
        c = ws.cell(r, j); c.border = THIN
        if c.font != INP: c.font = TXT
    if x.vys_mes < -8000:
        for j in range(1, 21): ws.cell(r, j).fill = RED if ws.cell(r, j).fill != YEL else YEL
    r += 1
last = r - 1
ws.cell(r, 1, 'Celkem').font = BOLD
ws.cell(r, 5, f'=SUM(E5:E{last})').number_format = NUM
for j in range(9, 18):
    col = get_column_letter(j); c = ws.cell(r, j, f'=SUM({col}5:{col}{last})'); c.font = BOLD; c.number_format = NUM
ws.cell(r, 18, f'=IFERROR(Q{r}/E{r}*30.4,0)').number_format = NUM; ws.cell(r, 18).font = BOLD
ws.cell(r, 2, f'=COUNTA(A5:A{last})&" aut"').font = BOLD
ws.cell(r + 2, 1, 'Červeně = ztráta víc než 8 tis. Kč za měsíc. Žlutě = náklady GFAR podle zadání (v datech chybí leasing). '
        'Doporučení je orientační podle podílu dní. Bratislava: podíly dní chybí (SK Renteon), dny = doba ve flotile podle CZ Renteonu.').font = NOTE
ws.freeze_panes = 'C5'; ws.auto_filter.ref = f'A4:T{last}'

# --- auta mimo Renteon ---
mimo = d[d.sipp.isna() & ~d.spz.str.startswith('(') & ~d.spz.str.contains('GLOBAL', na=False)]
m = mimo.groupby('spz').agg(model=('model', 'first'), pobocka=('pobocka', 'last'), vyn=('vynos_celkem', 'sum'), nak=('naklady_celkem', 'sum'),
                            pozn=('poznamka', lambda s: ', '.join(sorted(set(s.dropna()))))).reset_index()
m = m[(m.vyn - m.nak) < 0].sort_values(by='vyn')
w2 = wb.create_sheet('Mimo Renteon')
w2['A1'] = 'Auta mimo Renteon se ztrátou 2025–2026 (jen Vyfakturuj / HAVEX)'; w2['A1'].font = TITLE
w2['A2'] = 'Bez dnů provozu – nedá se přepočíst na měsíc. Často auta HAVEX/Fábián nebo vrácená auta, kde výnos chybí kvůli přiřazení SPZ.'; w2['A2'].font = NOTE
for i, (t, wd) in enumerate(zip(['SPZ', 'Model', 'Pobočka', 'Výnos', 'Náklady', 'Výsledek', 'Poznámka'], [10, 26, 8, 12, 12, 12, 30]), 1):
    c = w2.cell(4, i, t); c.font = H; c.fill = HF; w2.column_dimensions[get_column_letter(i)].width = wd
r = 5
for x in m.sort_values(by='vyn').itertuples(index=False):
    for j, v in enumerate([x.spz, x.model if isinstance(x.model, str) else '', x.pobocka], 1): w2.cell(r, j, v).font = INP
    for j, v in zip([4, 5], [x.vyn, x.nak]):
        c = w2.cell(r, j, int(v)); c.font = INP; c.number_format = NUM
    w2.cell(r, 6, f'=D{r}-E{r}').number_format = NUM
    w2.cell(r, 7, x.pozn).font = TXT
    r += 1
w2.cell(r, 1, 'Celkem').font = BOLD
for j in [4, 5, 6]:
    col = get_column_letter(j); c = w2.cell(r, j, f'=SUM({col}5:{col}{r-1})'); c.font = BOLD; c.number_format = NUM
w2.freeze_panes = 'B5'

out = B / 'output' / 'Ztratova_auta.xlsx'; wb.save(out)
print(out, len(a), 'aut v Renteonu,', len(m), 'mimo')
pd.set_option('display.width', 220)
print(a[['spz', 'model', 'sipp', 'pobocka', 'dny', 'kr', 'dl', 'vys', 'vys_mes']].round(0).to_string(index=False))
print('mimo celkem', round((m.vyn - m.nak).sum()))
