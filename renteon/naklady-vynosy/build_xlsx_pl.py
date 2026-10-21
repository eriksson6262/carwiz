import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
F=Font(name='Arial',size=10); B=Font(name='Arial',size=10,bold=True); BL=Font(name='Arial',size=10,color='0000FF'); G=Font(name='Arial',size=10,color='008000')
H=PatternFill('solid',fgColor='DDE4EE'); NF='#,##0;(#,##0);-'
def num(v):
    try: return float(v)
    except: return v
def sheet(wb,name,rows,hdr):
    ws=wb.create_sheet(name); ws.append(hdr)
    for r in rows: ws.append([num(r[k]) if k not in('rok','mesic','datum','misto','produkt','karta4','druh','mena','odberatel','dodavatel','pobocka','spz','model','sipp','pobocka','poznamka','doklad','obdobi','vin','kategorie','zdroj_email_id') else r[k] for k in hdr])
    for c in ws[1]: c.font=B; c.fill=H
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=BL if isinstance(c.value,(int,float)) else F
            if isinstance(c.value,(int,float)): c.number_format=NF
    for i,h in enumerate(hdr,1): ws.column_dimensions[L(i)].width=max(10,len(h)+2)
    ws.freeze_panes='A2'; return ws
wb=Workbook(); S=wb.active; S.title='Souhrn'
pl=list(csv.DictReader(open('pl_pobocky.csv'))); hp=list(pl[0].keys())
D=sheet(wb,'Data',pl,hp); n=len(pl)+1
rc=len(hp)+1; D.cell(1,rc,'rok').font=B; D.cell(1,rc).fill=H
for i in range(2,n+1): D.cell(i,rc,f'=LEFT(A{i},4)').font=F
col={h:L(i+1) for i,h in enumerate(hp)}; RK=L(rc)
S['A1']='Náklady vs výnosy půjčovny 01/2024–10/2026 (Kč bez DPH)'; S['A1'].font=Font(name='Arial',size=12,bold=True)
br=['PRG','OSR','BTS','HAVEX_ZF']; names={'PRG':'Praha','OSR':'Ostrava','BTS':'Bratislava','HAVEX_ZF':'HAVEX zpětná fakt.'}
items=[('Výnos Renteon','vynos_renteon'),('Výnos Vyfakturuj / HAVEX','vynos_vyfakturuj'),('Výnosy celkem','vynos_celkem'),('Leasing (kalendáře)','leasing'),('Leasing (odhad)','leasing_odhad'),('Carnetto – půjčovné EFS','carnetto_leasing'),('Pojištění','pojisteni'),('Carnetto – pojištění ČGP','carnetto_pojisteni'),('Servis','servis'),('Opravy','opravy'),('Ostatní (mytí, známky)','ostatni'),('Palivo CCS','palivo'),('Nájem pobočky (faktury)','najem_pobocky'),('Nájem pobočky (odhad)','najem_odhad'),('Náklady celkem','naklady_celkem'),('Výsledek','vysledek')]
r=2
for rok in ['2024','2025','2026','Celkem']:
    r+=1; S.cell(r,1,rok if rok!='2026' else '2026 (I–X)').font=Font(name='Arial',size=11,bold=True)
    r+=1; hdr=['Položka']+[names[b] for b in br]+['Celkem']
    for j2,v in enumerate(hdr,1): c=S.cell(r,j2,v); c.font=B; c.fill=H
    top=r+1
    for t,k in items:
        r+=1; S.cell(r,1,t).font=B if k in('vynos_celkem','naklady_celkem','vysledek') else F
        for j2,b in enumerate(br,2):
            f=f'Data!${col[k]}$2:${col[k]}${n},Data!$B$2:$B${n},"{b}"'
            if rok!='Celkem': f+=f',Data!${RK}$2:${RK}${n},"{rok}"'
            c=S.cell(r,j2,f"=SUMIFS({f})"); c.font=G; c.number_format=NF
        c=S.cell(r,6,f"=SUM(B{r}:E{r})"); c.number_format=NF; c.font=B
    vr=top+2; r+=1; S.cell(r,1,'Marže').font=B
    for j2 in range(2,7):
        c=S.cell(r,j2,f"=IF({L(j2)}{vr}=0,0,{L(j2)}{r-1}/{L(j2)}{vr})"); c.number_format='0.0%;(0.0%);-'; c.font=F
    r+=1
S.column_dimensions['A'].width=28
for j in range(2,7): S.column_dimensions[L(j)].width=16
au=list(csv.DictReader(open('pl_auta.csv'))); sheet(wb,'Auta',au,list(au[0].keys()))
cm=list(csv.DictReader(open('carnetto_mesice.csv'))); sheet(wb,'Carnetto',cm,list(cm[0].keys()))
cd=list(csv.DictReader(open('carnetto.csv'))); sheet(wb,'Carnetto detail',cd,list(cd[0].keys()))
for nm,fn in [('Nájmy','najmy.csv'),('Palivo','palivo.csv'),('Servis Rejdík SIMO AP','servis2.csv'),('Servis 2024-2025','servis_2024.csv')]:
    rr=list(csv.DictReader(open(fn))); sheet(wb,nm,rr,list(rr[0].keys()))
P=wb.create_sheet('Poznámky')
for t in ['Zdroje: Renteon (výnosy po pobočkách a vozech), Vyfakturuj RX Auto + RX Auto SK, Spark (HAVEX kalendáře, EFS, Kooperativa, Generali, D.A.S., servisy).',
 'Carnetto RX: půjčovné EURO FLEET SERVIS (11/2025–09/2026, 11 faktur) a pojištění flotily ČGP 6287974108 (04–09/2026) – přefakturováno na RX Auto, rozpis po SPZ z faktur/vyúčtování.',
 'Faktury EFS i vyúčtování ČGP za 10/2026 zatím nepřišly – říjen je v nákladech podhodnocen.',
 'Bratislava přepočtena kurzem 24,3 Kč/€ (předpoklad).',
  'HAVEX_ZF = vozy se zpětnou fakturací HAVEX Mobility: výnos z Vyfakturuj vs leasing těch vozů.',
 'Chybí: mzdy, provize brokerů, franšíza, IT. Servisní faktury ve schránce až od 09/2025.',
 'Nájmy: CTP (letiště Praha, část v EUR ×24,3), BTS aero (RX Auto SK, EUR), Letiště Ostrava. Do 06/2025 faktury většinou v poště nejsou → odhad základního nájmu: PRG 1 237 €/měs (01/2024–06/2025), BTS 1 570,25 €/měs (05/2024–06/2025, 08/2025), OSR 24 738 Kč/měs (01/2024–09/2025). Obratové nájemné 2024 chybí.',
 'Leasing do 06/2025 = odhad z HAVEX přehledu zápůjček (měsíční cena smlouvy pro rata), platební kalendáře v poště až od 07/2025.',
 'Pojištění: Kooperativa 2024–2026, ČGP RS 4785635732 (chybí 01/2024, 01–02/2025), D.A.S. od 05/2025.',
 'Servisy 2024–08/2025: HAVEX přefakturace, Rejdík, Auto Palace, SIMO CAR PLUS aj. – část faktur v poště chybí (Rejdík č. 21–48, SIMO 08/2024–06/2025, mytí AP před 11/2024).',
 'Palivo: měsíční doklady CCS (5 karet, bez SPZ → pobočka podle místa čerpání); tankování v SK včetně slovenské DPH (neodpočitatelná). 10/2026 zatím chybí.',
 'Modrá = vstupní data, zelená = odkaz na list Data, černá = výpočet.']:
    P.append([t]); P.cell(P.max_row,1).font=F
P.column_dimensions['A'].width=140
wb.save('Naklady_vs_vynosy_2024-2026.xlsx')
