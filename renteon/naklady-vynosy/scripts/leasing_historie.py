#!/usr/bin/env python3
"""Leasing HAVEX po měsících a SPZ z přehledu zápůjček (list „Vozy“), celá historie od 2020.
Vstup: data/zdroje/prehled_zapujcek_*.xlsx (poslední od M. Dvořákové). Výstup: data/leasing_historie.csv
(mesic, spz, castka_bez_dph, typ [napřímo | zpětná fakturace], zpet_mesic = částka k zpětné fakturaci na HAVEX, smlouva).
Splátka = cena_najmu_mesic × dny v měsíci / dny měsíce (stejně jako odhad v leasing_build.py); poznámka „ukončeno k d.m.rrrr“ zkracuje konec."""
import sys, re, glob, calendar, datetime as dt, pathlib
import pandas as pd
from openpyxl import load_workbook
B = pathlib.Path(__file__).resolve().parents[1]
f = sys.argv[1] if len(sys.argv) > 1 else sorted(glob.glob(str(B / 'data/zdroje/prehled_zapujcek_*.xlsx')))[-1]
ws = load_workbook(f, data_only=True)['Vozy']
OUT_DO = dt.date(2026, 10, 31)
rows = []
for r in ws.iter_rows(min_row=4, max_col=13, values_only=True):
    zac, kon, sml, fin, firma, auto, cena, marze, km, zpet, prist, najeto, pozn = r
    if not isinstance(zac, (dt.date, dt.datetime)) or not cena:
        continue
    m = re.search(r'rz:\s*([0-9A-Z]{3}\s?[0-9A-Z]{4}|[A-Z]{2}\d{3}[A-Z]{2})', str(auto or ''), re.I)
    spz = m.group(1).replace(' ', '').upper() if m else ''
    z = zac.date() if isinstance(zac, dt.datetime) else zac
    k = (kon.date() if isinstance(kon, dt.datetime) else kon) if isinstance(kon, (dt.date, dt.datetime)) else OUT_DO
    mu = re.search(r'ukon[čc]en[oa]?\s*(?:dohodou\s*)?k\s*(\d{1,2})\.\s*(\d{1,2})\.\s*(\d{4})', str(pozn or ''), re.I)
    if mu:
        k = min(k, dt.date(int(mu.group(3)), int(mu.group(2)), int(mu.group(1))))
    k = min(k, OUT_DO)
    try:
        cena = float(cena); zp = float(zpet or 0)
    except (TypeError, ValueError):
        continue
    typ = 'zpětná fakturace' if zp else 'napřímo'
    y, mo = z.year, z.month
    while dt.date(y, mo, 1) <= k:
        dim = calendar.monthrange(y, mo)[1]
        s, e = max(z, dt.date(y, mo, 1)), min(k, dt.date(y, mo, dim))
        days = (e - s).days + 1
        if days > 0:
            rows.append(dict(mesic=f'{y}-{mo:02d}', spz=spz, castka_bez_dph=round(cena * days / dim, 2), typ=typ,
                             zpet_mesic=round(zp * days / dim, 2), smlouva=str(sml), financovani=fin or ''))
        mo += 1
        if mo == 13:
            y, mo = y + 1, 1
df = pd.DataFrame(rows)
df.to_csv(B / 'data' / 'leasing_historie.csv', index=False)
print(f, len(df), 'řádků')
print(df.groupby([df.mesic.str[:4], 'typ']).castka_bez_dph.sum().round(-3).unstack())
