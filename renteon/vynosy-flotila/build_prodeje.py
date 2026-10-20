import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter as L
D=sys.argv[1] if len(sys.argv)>1 else '.'
F='Arial'; hdr=Font(name=F,bold=True,color='FFFFFF'); hfill=PatternFill('solid',fgColor='1F2937')
f=Font(name=F); fb=Font(name=F,bold=True); blue=Font(name=F,color='0000FF'); green=Font(name=F,color='008000'); grey=Font(name=F,italic=True,color='6B7280')
YC={'2024':'2BA58F','2025':'E07B39','2026':'2F6DB5'}
mes=list(csv.DictReader(open(f'{D}/prodeje_mesice.csv'))); cars=list(csv.DictReader(open(f'{D}/prodeje_auta.csv'))); odb=list(csv.DictReader(open(f'{D}/prodeje_odberatele.csv')))
wb=Workbook()
# Měsíce
ws=wb.active; ws.title='Měsíce'
cols=[('mesic','Měsíc'),('zeme','Firma'),('mena','Měna'),('trzby_faktury','Tržby – faktury'),('trzby_danove_doklady','Tržby – daňové doklady'),('pocet_aut','Prodaných aut'),('pocet_dokladu','Dokladů')]
ws.append([h for _,h in cols]+['Tržby celkem'])
for r in mes: ws.append([r[k] if i<3 else int(float(r[k])) for i,(k,_) in enumerate(cols)])
N=len(mes)+1
for i in range(2,N+1): ws.cell(i,8,f'=D{i}+E{i}')
for c in ws[1]: c.font=hdr; c.fill=hfill; c.alignment=Alignment(wrap_text=True,horizontal='center')
for row in ws.iter_rows(min_row=2,max_row=N):
    for c in row: c.font=blue if 4<=c.column<=7 else f; c.number_format='#,##0' if c.column>=4 else 'General'
for i,w in enumerate([10,9,7,15,18,12,10,15],1): ws.column_dimensions[L(i)].width=w
ws.freeze_panes='A2'
ws.cell(N+2,1,'Zdroj: Vyfakturuj, RX Auto s.r.o. (CZ, Kč) a RX Auto SK s. r. o. (SK, €). Jen faktury a daňové doklady k přijatým zálohám za prodej vozidel, podle DUZP, bez DPH, bez stornovaných. Konečná faktura odečítá zálohu, součet je čistý.').font=grey
ws.cell(N+3,1,'Počet aut = unikátní VIN; auto se započítá v měsíci prvního dokladu (u záloh daňový doklad). Pronájmy, jistiny a HAVEX Mobility vyloučeny. Říjen 2026 neúplný (do 20. 10.).').font=grey
R=lambda c: f"'Měsíce'!${c}$2:${c}${N}"
# Souhrn
ss=wb.create_sheet('Souhrn',0)
ss['A1']='Prodeje aut 2024–2026'; ss['A1'].font=Font(name=F,bold=True,size=14)
ss['A2']='Tržby bez DPH z faktur a daňových dokladů (Vyfakturuj). Hodnoty počítané z listu Měsíce.'; ss['A2'].font=grey
heads=['Firma','Rok','Tržby celý rok','Prodaných aut','Ø na auto','Tržby I–IX','Aut I–IX','Změna tržeb I–IX','Změna aut I–IX']
for i,h in enumerate(heads,1): c=ss.cell(4,i,h); c.font=hdr; c.fill=hfill; c.alignment=Alignment(wrap_text=True,horizontal='center')
ss.row_dimensions[4].height=30
r=5
for z,name,cur in [('CZ','RX Auto s.r.o.','Kč'),('SK','RX Auto SK s. r. o.','€')]:
    first=r
    for y in ['2024','2025','2026']:
        ss.cell(r,1,f'{name} ({cur})').font=f; c=ss.cell(r,2,y); c.font=Font(name=F,bold=True,color=YC[y])
        ss.cell(r,3,f'=SUMIFS({R("H")},{R("B")},$A$9999&"{z}",{R("A")},B{r}&"*")'.replace('$A$9999&','')).number_format='#,##0'
        ss.cell(r,4,f'=SUMIFS({R("F")},{R("B")},"{z}",{R("A")},B{r}&"*")').number_format='#,##0'
        ss.cell(r,5,f'=IFERROR(C{r}/D{r},"")').number_format='#,##0'
        ss.cell(r,6,f'=SUMPRODUCT(({R("B")}="{z}")*(LEFT({R("A")},4)=B{r})*(VALUE(RIGHT({R("A")},2))<=9)*{R("H")})').number_format='#,##0'
        ss.cell(r,7,f'=SUMPRODUCT(({R("B")}="{z}")*(LEFT({R("A")},4)=B{r})*(VALUE(RIGHT({R("A")},2))<=9)*{R("F")})').number_format='#,##0'
        if r>first:
            ss.cell(r,8,f'=IFERROR(F{r}/F{r-1}-1,"")').number_format='+0.0%;-0.0%;0.0%'
            ss.cell(r,9,f'=IFERROR(G{r}/G{r-1}-1,"")').number_format='+0.0%;-0.0%;0.0%'
        for c in range(3,10): ss.cell(r,c).font=f
        r+=1
    r+=1
# fix tržby formula (simple)
for rr in range(5,r):
    if ss.cell(rr,2).value: 
        z='CZ' if 'SK' not in str(ss.cell(rr,1).value) else 'SK'
        ss.cell(rr,3,f'=SUMIFS({R("H")},{R("B")},"{z}",{R("A")},B{rr}&"*")')
for i,w in enumerate([24,8,16,14,14,16,10,14,12],1): ss.column_dimensions[L(i)].width=w
# monthly CZ table for chart
t0=r+2
ss.cell(t0-1,1,'RX Auto s.r.o. po měsících (Kč)').font=Font(name=F,bold=True,size=12)
for i,h in enumerate(['Měsíc','Faktury','Daňové doklady','Tržby celkem','Prodaných aut'],1): c=ss.cell(t0,i,h); c.font=hdr; c.fill=hfill
months=[x['mesic'] for x in mes if x['zeme']=='CZ']
for j,mo in enumerate(months):
    rr=t0+1+j; ss.cell(rr,1,mo).font=f
    for ci,col in [(2,'D'),(3,'E'),(4,'H'),(5,'F')]:
        c=ss.cell(rr,ci,f'=SUMIFS({R(col)},{R("A")},$A{rr},{R("B")},"CZ")'); c.number_format='#,##0'; c.font=green
last=t0+len(months)
bar=BarChart(); bar.type='col'; bar.grouping='stacked'; bar.overlap=100; bar.gapWidth=40; bar.title='RX Auto s.r.o.: tržby z prodeje aut (Kč) a počet aut'
bar.add_data(Reference(ss,min_col=2,max_col=3,min_row=t0,max_row=last),titles_from_data=True); bar.set_categories(Reference(ss,min_col=1,min_row=t0+1,max_row=last))
for s,col in zip(bar.series,['2F6DB5','9DB8E0']): s.graphicalProperties.solidFill=col; s.graphicalProperties.line.solidFill=col
ln=LineChart(); ln.add_data(Reference(ss,min_col=5,min_row=t0,max_row=last),titles_from_data=True)
s=ln.series[0]; s.graphicalProperties.line.solidFill='1F2937'; s.graphicalProperties.line.width=25000; s.marker.symbol='circle'; s.marker.size=5; s.marker.graphicalProperties.solidFill='1F2937'
ln.y_axis.axId=200; ln.y_axis.title='Počet aut'; ln.y_axis.crosses='max'; ln.y_axis.majorGridlines=None
bar+=ln; bar.height=11; bar.width=30; bar.legend.position='b'; bar.y_axis.title='Kč'
ss.add_chart(bar,f'G{t0}')
# Auta
wa=wb.create_sheet('Auta')
ac=[('datum','Datum'),('zeme','Firma'),('vuz','Vůz'),('vin','VIN'),('odberatel','Odběratel'),('castka','Částka bez DPH'),('mena','Měna'),('doklad','Doklad')]
wa.append([h for _,h in ac])
for r_ in sorted(cars,key=lambda x:x['datum']): wa.append([int(r_[k]) if k=='castka' else r_[k] for k,_ in ac])
for c in wa[1]: c.font=hdr; c.fill=hfill
for row in wa.iter_rows(min_row=2):
    for c in row: c.font=f
    row[5].number_format='#,##0'
for i,w in enumerate([11,7,55,20,34,14,6,14],1): wa.column_dimensions[L(i)].width=w
wa.freeze_panes='A2'; wa.auto_filter.ref=f'A1:H{wa.max_row}'
# Odběratelé
wo=wb.create_sheet('Odběratelé')
wo.append(['Rok','Firma','Odběratel','Tržby bez DPH','Měna'])
for r_ in sorted(odb,key=lambda x:(x['rok'],-float(x['trzby']))): wo.append([r_['rok'],r_['zeme'],r_['odberatel'],int(float(r_['trzby'])),r_['mena']])
for c in wo[1]: c.font=hdr; c.fill=hfill
for row in wo.iter_rows(min_row=2):
    for c in row: c.font=f
    row[3].number_format='#,##0'
for i,w in enumerate([7,7,46,16,6],1): wo.column_dimensions[L(i)].width=w
wo.freeze_panes='A2'; wo.auto_filter.ref=f'A1:E{wo.max_row}'
wb.save(sys.argv[2] if len(sys.argv)>2 else 'Prodeje_aut_2024-2026.xlsx')
