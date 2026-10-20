import sys
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter as L

rows=list(csv.DictReader(open(sys.argv[1] if len(sys.argv)>1 else 'data/vynosy_flotila.csv')))
F='Arial'; C24='2BA58F'; C25='E07B39'; C26='2F6DB5'
YEARS=['2024','2025','2026']; YC={'2024':C24,'2025':C25,'2026':C26}
hdr=Font(name=F,bold=True,color='FFFFFF'); hfill=PatternFill('solid',fgColor='1F2937')
f=Font(name=F); fb=Font(name=F,bold=True); blue=Font(name=F,color='0000FF'); grey=Font(name=F,italic=True,color='6B7280')
thin=Border(bottom=Side(style='thin',color='D1D5DB'))
wb=Workbook()

# ---- Data
ds=wb.active; ds.title='Data'
dcols=[('mesic','Měsíc'),('pobocka','Pobočka'),('zeme','Země'),('mena','Měna'),('vynos_celkem','Renteon výnos celkem bez DPH'),
 ('najem','Nájem'),('pojisteni','Pojištění'),('poplatky_bez_pokut','Poplatky bez pokut'),('skody','Škody'),('ostatni','Ostatní'),
 ('pokuty','Pokuty (přeúčtované)'),('zprac_pokut','Zpracování pokut'),('pocet_pokut','Počet pokut'),('smluv','Smluv'),('dni','Dní'),('vozidel_renteon','Ø vozidel Renteon'),
 ('vf_najem_renteon_auta','Vyfakturuj – auta z Renteonu'),('vf_najem_mimo_renteon','Vyfakturuj – auta mimo Renteon'),('vf_pokuty','Vyfakturuj – pokuty'),('vf_vozidel_mimo','Vozidel mimo Renteon')]
NI=len(dcols)
ds.append([h for _,h in dcols]+['Výnos bez pokut (Renteon + Vyfakturuj)','Ø vozidel celkem','Vyfakturuj celkem','Pokuty vč. zprac. celkem'])
for r in rows:
    ds.append([r[k] if i<4 else (float(r[k]) if k=='vozidel_renteon' else int(float(r[k]))) for i,(k,_) in enumerate(dcols)])
N=len(rows)+1
CL={k:L(i+1) for i,(k,_) in enumerate(dcols)}
for j,k in enumerate(['vynos_bez_pokut','vozidel','vf','pok_all']): CL[k]=L(NI+1+j)
for r in range(2,N+1):
    ds[f"{CL['vynos_bez_pokut']}{r}"]=f'={CL["vynos_celkem"]}{r}-{CL["pokuty"]}{r}-{CL["zprac_pokut"]}{r}+{CL["vf_najem_renteon_auta"]}{r}+{CL["vf_najem_mimo_renteon"]}{r}'
    ds[f"{CL['vozidel']}{r}"]=f'={CL["vozidel_renteon"]}{r}+{CL["vf_vozidel_mimo"]}{r}'
    ds[f"{CL['vf']}{r}"]=f'={CL["vf_najem_renteon_auta"]}{r}+{CL["vf_najem_mimo_renteon"]}{r}'
    ds[f"{CL['pok_all']}{r}"]=f'={CL["pokuty"]}{r}+{CL["zprac_pokut"]}{r}+{CL["vf_pokuty"]}{r}'
for c in ds[1]: c.font=hdr; c.fill=hfill; c.alignment=Alignment(wrap_text=True,horizontal='center')
ds.row_dimensions[1].height=44
for row in ds.iter_rows(min_row=2,max_row=N):
    for c in row:
        c.font=blue if 5<=c.column<=NI else f
        if c.column>=5: c.number_format='#,##0'
    row[15].number_format='0.0'; row[NI+1].number_format='0.0'
for i in range(1,NI+5): ds.column_dimensions[L(i)].width=10 if i<=4 else 13
ds.freeze_panes='C2'
notes=['Zdroj Renteon: CZ (rxauto.s3) a SK (cwslovakia.s6), report Realizace po kancelářích – podle data vyúčtování, bez DPH; CZ v Kč, SK v €.',
 'Pokuty = položky dokladů typu dopravní pokuta / mýtné / parkování (typ služby 14) a poplatek za zpracování, podle data dokladu; přiřazeno pobočce rezervace, doklady bez smlouvy pobočce dokladu. Odečteny z poplatků.',
 'Zdroj Vyfakturuj: RX Auto s.r.o. (PRG, OSR) a RX Auto SK s. r. o. (BTS), faktury, daňové, opravné a příjmové doklady podle DUZP, bez DPH, bez stornovaných; bez HAVEX Mobility; bez prodejů aut (viz samostatný přehled) a jistin. Položky s přestupkem/pokutou/admin. poplatkem = Vyfakturuj – pokuty.',
 'Pobočka pro Vyfakturuj: podle SPZ v textu položky a toho, na které pobočce (OUT) byl vůz v daném měsíci v Renteonu nejvíc dní; GLOBAL ASSISTANCE = Ostrava; auta, která v Renteonu nejsou (dlouhodobá flotila), = Praha a jsou přičtena k počtu vozidel Prahy.',
 'Ø vozidel Renteon: report Stav flotily po kancelářích, flotila Pronájem vozového parku, průměr v měsíci. Staženo 20.–21. 10. 2026; říjen 2026 neúplný.']
for i,t in enumerate(notes): ds.cell(N+2+i,1,t).font=grey
D=lambda col: f"Data!${CL[col]}$2:${CL[col]}${N}"

MES=['leden','únor','březen','duben','květen','červen','červenec','srpen','září','říjen','listopad','prosinec']
names={'PRG':'Praha (PRG)','OSR':'Ostrava (OSR)','BTS':'Bratislava (BTS)'}
cur={'PRG':'Kč','OSR':'Kč','BTS':'€'}
fields=[('vynos_bez_pokut','Výnos bez pokut'),('najem','Nájem Renteon'),('pojisteni','Pojištění'),('poplatky_bez_pokut','Poplatky bez pokut'),
        ('skody','Škody'),('vf','Vyfakturuj'),('pok_all','Pokuty vč. zprac.'),('smluv','Smluv Renteon'),('vozidel','Ø vozidel'),('nv','Výnos na vůz')]
NF=len(fields); IV=0; IVOZ=8; ISML=7

def cond(key,p,expr):
    return f'=IF(COUNTIFS({D("mesic")},{key},{D("pobocka")},"{p}")=0,"",{expr})'
def sif(col,key,p): return f'SUMIFS({D(col)},{D("mesic")},{key},{D("pobocka")},"{p}")'

for p in ['PRG','OSR','BTS']:
    ws=wb.create_sheet(p)
    ws['A1']=f"{names[p]} – výnos bez DPH a bez pokut ({cur[p]}), smlouvy a flotila po měsících"; ws['A1'].font=Font(name=F,bold=True,size=14)
    for col,t in [(1,'Měsíc'),(2,'Klíč 2024'),(3,'Klíč 2025'),(4,'Klíč 2026')]:
        c=ws.cell(3,col,t); c.font=hdr; c.fill=hfill
    start=5
    for yi,y in enumerate(YEARS):
        for fi,(k,t) in enumerate(fields):
            c=ws.cell(3,start+yi*NF+fi,f"{t} {y}"); c.font=hdr; c.fill=PatternFill('solid',fgColor=YC[y])
            c.alignment=Alignment(wrap_text=True,horizontal='center')
    for m in range(12):
        r=4+m
        ws.cell(r,1,MES[m]).font=f
        for yi,y in enumerate(YEARS): ws.cell(r,2+yi,f"{y}-{m+1:02d}")
        for yi in range(3):
            key=f"${L(2+yi)}{r}"; base=start+yi*NF
            for fi,(k,t) in enumerate(fields):
                c=ws.cell(r,base+fi)
                if k=='nv':
                    v=L(base+IV); z=L(base+IVOZ); c.value=f'=IF(OR({v}{r}="",{z}{r}=""),"",{v}{r}/{z}{r})'
                else:
                    c.value=cond(key,p,sif(k,key,p))
                c.font=f; c.number_format='0.0' if k=='vozidel' else '#,##0'; c.border=thin
    tr=17
    ws.cell(tr,1,'Leden–září').font=fb
    for yi in range(3):
        base=start+yi*NF
        for fi,(k,t) in enumerate(fields):
            col=L(base+fi); c=ws.cell(tr,base+fi)
            if k=='vozidel': c.value=f'=AVERAGE({col}4:{col}12)'
            elif k=='nv': c.value=f'={L(base+IV)}{tr}/{L(base+IVOZ)}{tr}/9'
            else: c.value=f'=SUM({col}4:{col}12)'
            c.font=fb; c.number_format='0.0' if k=='vozidel' else '#,##0'
    ws.cell(18,1,'Změna proti předch. roku').font=fb
    for yi in (1,2):
        for fi in range(NF):
            c=ws.cell(18,start+yi*NF+fi,f'=IFERROR({L(start+yi*NF+fi)}17/{L(start+(yi-1)*NF+fi)}17-1,"")'); c.number_format='+0.0%;-0.0%;0.0%'; c.font=fb
    ws.cell(19,1,'Výnos bez pokut = Renteon (bez pokut a jejich zpracování) + Vyfakturuj (dlouhodobé pronájmy, bez HAVEX Mobility). Ø vozidel = Renteon + auta mimo Renteon fakturovaná ve Vyfakturuj. Výnos na vůz = výnos / Ø vozidel. Říjen 2026 neúplný.').font=grey
    ws.column_dimensions['A'].width=16
    for cc in 'BCD': ws.column_dimensions[cc].hidden=True
    for i in range(5,5+3*NF): ws.column_dimensions[L(i)].width=12
    ws.row_dimensions[3].height=32; ws.freeze_panes='E4'

    # chart 1: bars výnos + lines vozidla, colour per year
    bar=BarChart(); bar.type='col'; bar.grouping='clustered'; bar.overlap=-10; bar.gapWidth=60
    bar.title=f"{names[p]}: výnos bez pokut ({cur[p]}), Ø vozidel (plná) a smlouvy (přerušovaná)"; bar.y_axis.title=f"Výnos {cur[p]}"
    for yi,col in enumerate([C24,C25,C26]):
        bar.add_data(Reference(ws,min_col=start+yi*NF+IV,min_row=3,max_row=15),titles_from_data=True)
        s=bar.series[-1]; s.graphicalProperties.solidFill=col; s.graphicalProperties.line.solidFill=col
    bar.set_categories(Reference(ws,min_col=1,min_row=4,max_row=15))
    line=LineChart()
    for yi,col in enumerate([C24,C25,C26]):
        for ci,dash_ in ((IVOZ,None),(ISML,'dash')):
            line.add_data(Reference(ws,min_col=start+yi*NF+ci,min_row=3,max_row=13 if yi==2 else 15),titles_from_data=True)
            s=line.series[-1]; s.graphicalProperties.line.solidFill=col; s.graphicalProperties.line.width=32000 if dash_ is None else 22000
            if dash_: s.graphicalProperties.line.dashStyle='dash'
            s.marker.symbol='circle' if dash_ is None else 'triangle'; s.marker.size=6; s.marker.graphicalProperties.solidFill=col; s.marker.graphicalProperties.line.solidFill=col; s.smooth=False
    line.y_axis.axId=200; line.y_axis.title='Ø vozidel / počet smluv'; line.y_axis.crosses='max'; line.y_axis.majorGridlines=None
    bar+=line; bar.height=11; bar.width=30; bar.legend.position='b'
    ws.add_chart(bar,'A21')

    # chart 2: smlouvy po měsících (same colours)
    sm=BarChart(); sm.type='col'; sm.grouping='clustered'; sm.gapWidth=60; sm.title=f"{names[p]}: počet smluv"
    for yi,col in enumerate([C24,C25,C26]):
        sm.add_data(Reference(ws,min_col=start+yi*NF+ISML,min_row=3,max_row=15),titles_from_data=True)
        s=sm.series[-1]; s.graphicalProperties.solidFill=col; s.graphicalProperties.line.solidFill=col
    sm.set_categories(Reference(ws,min_col=1,min_row=4,max_row=15)); sm.height=10; sm.width=18; sm.legend.position='b'
    ws.add_chart(sm,'S21')

    # timeline breakdown
    t0=45
    ws.cell(t0-1,1,'Rozpad výnosu po měsících').font=Font(name=F,bold=True,size=12)
    th=['Měsíc','Nájem','Pojištění','Poplatky bez pokut','Škody','Ostatní','Vyfakturuj – auta z Renteonu','Vyfakturuj – auta mimo Renteon','Pokuty vč. zprac.','Výnos celkem vč. pokut','Smluv']
    for i,h in enumerate(th): c=ws.cell(t0,1+i,h); c.font=hdr; c.fill=hfill
    months=sorted({r['mesic'] for r in rows})
    for j,mk in enumerate(months):
        r=t0+1+j; ws.cell(r,1,mk).font=f; key=f'$A{r}'
        for i,k in enumerate(['najem','pojisteni','poplatky_bez_pokut','skody','ostatni','vf_najem_renteon_auta','vf_najem_mimo_renteon','pok_all']):
            c=ws.cell(r,2+i,'='+sif(k,key,p)); c.number_format='#,##0'; c.font=f
        c=ws.cell(r,10,f'=SUM(B{r}:I{r})'); c.number_format='#,##0'; c.font=f
        c=ws.cell(r,11,'='+sif('smluv',key,p)); c.number_format='#,##0'; c.font=f
    last=t0+len(months)
    st=BarChart(); st.type='col'; st.grouping='stacked'; st.overlap=100; st.gapWidth=40
    st.title=f"{names[p]}: rozpad výnosu – Renteon, Vyfakturuj a vyčleněné pokuty ({cur[p]})"
    st.add_data(Reference(ws,min_col=2,max_col=9,min_row=t0,max_row=last),titles_from_data=True)
    st.set_categories(Reference(ws,min_col=1,min_row=t0+1,max_row=last))
    for s,col in zip(st.series,['2F6DB5','2BA58F','E0B33A','D2555A','8B5CF6','0EA5E9','14B8A6','9CA3AF']):
        s.graphicalProperties.solidFill=col; s.graphicalProperties.line.solidFill=col
    st.height=10; st.width=26; st.legend.position='b'
    ws.add_chart(st,f'M{t0}')

# ---- Souhrn
ss=wb.create_sheet('Souhrn',0)
ss['A1']='Výnosy, smlouvy a flotila – leden–září 2024, 2025 a 2026'; ss['A1'].font=Font(name=F,bold=True,size=14)
ss['A2']='Výnos bez DPH a bez přeúčtovaných pokut: Renteon + Vyfakturuj (bez HAVEX Mobility, bez prodejů aut). CZ v Kč, SK v €. Pokuty vyčleněny. Bratislava od března 2024.'; ss['A2'].font=grey
metrics=[('vynos_bez_pokut','Výnos bez pokut (Renteon + Vyfakturuj)','#,##0'),('renf','z toho Renteon','#,##0'),('vf','z toho Vyfakturuj','#,##0'),('smluv','Smluv Renteon','#,##0'),('vozidel','Ø vozidel','0.0'),('nv','Výnos na vůz/měs.','#,##0'),('vs','Renteon výnos na smlouvu','#,##0'),
         ('najem','Nájem Renteon','#,##0'),('pojisteni','Pojištění','#,##0'),('poplatky_bez_pokut','Poplatky bez pokut','#,##0'),('skody','Škody','#,##0'),('pok_all','Pokuty vč. zprac.','#,##0')]
fidx={k:i for i,(k,_) in enumerate(fields)}
r=4
for p in ['PRG','OSR','BTS']:
    ss.cell(r,1,f"{names[p]} ({cur[p]})").font=Font(name=F,bold=True,size=12)
    for i,t in enumerate(['Ukazatel','2024','2025','2026','2025 vs 2024','2026 vs 2025']):
        c=ss.cell(r+1,1+i,t); c.font=hdr; c.fill=hfill if i in (0,4,5) else PatternFill('solid',fgColor=YC[t])
    for j,(k,t,fmt) in enumerate(metrics):
        rr=r+2+j; ss.cell(rr,1,t).font=f
        for yi in range(3):
            if k=='vs': v=f'=IFERROR({L(2+yi)}{r+3}/{L(2+yi)}{r+5},"")'; fc='000000'
            elif k=='renf': v=f'={L(2+yi)}{r+2}-{L(2+yi)}{r+4}'; fc='000000'
            else: v=f"={p}!{L(5+yi*NF+fidx[k])}17"; fc='008000'
            c=ss.cell(rr,2+yi,v); c.font=Font(name=F,color=fc); c.number_format=fmt
        for ci,(a,b) in enumerate([('C','B'),('D','C')]):
            c=ss.cell(rr,5+ci,f'=IFERROR({a}{rr}/{b}{rr}-1,"")'); c.number_format='+0.0%;-0.0%;0.0%'; c.font=f
    r+=2+len(metrics)+2
ss.column_dimensions['A'].width=36
for cc in 'BCDEF': ss.column_dimensions[cc].width=15
ss.cell(r,1,'Pozn.: do srpna 2025 se pokuty za celou flotilu fakturovaly převážně přes pobočku Ostrava (účty bez smlouvy); od září 2025 jsou u smluv pobočky rezervace.').font=grey
wb.save(sys.argv[2] if len(sys.argv)>2 else 'output/Vynosy_flotila.xlsx')
