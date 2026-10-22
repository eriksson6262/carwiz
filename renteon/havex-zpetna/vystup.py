#!/usr/bin/env python3
"""Excel „Nevyfakturováno HAVEX“ po vozidlech. Vstup: output/obdobi.csv, output/polozky_spz.csv (parovani.py), data/vozy.pkl,
data/jine_faktury.csv (faktury k autu na jiné odběratele než HAVEX – prodej, přímý pronájem klientovi)."""
import pathlib, re, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
B = pathlib.Path(__file__).parent
P = pd.read_csv(B / 'output/obdobi.csv', parse_dates=['od', 'do']); SI = pd.read_csv(B / 'output/polozky_spz.csv')
V = pd.read_pickle(B / 'data/vozy.pkl'); J = pd.read_csv(B / 'data/jine_faktury.csv', sep='|')
OD = pd.Timestamp('2021-09-01')
F = 'Arial'; H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3A5F'); T = Font(name=F); BO = Font(name=F, bold=True)
NOTE = Font(name=F, italic=True, color='666666', size=9); NUM = '#,##0;-#,##0;-'; THIN = Border(bottom=Side(style='thin', color='DDDDDD'))
RED = PatternFill('solid', fgColor='FDE2E1'); YEL = PatternFill('solid', fgColor='FFF4CC')
fmt = lambda v: f'{v:,.0f}'.replace(',', ' ')

def souhrn(df, si):
    e = df.groupby('spz').agg(ma=('castka', 'sum'), bez=('bez_spz', 'sum'), od=('od', 'min'), do=('do', 'max'),
        smlouvy=('smlouva', lambda s: ', '.join(dict.fromkeys(str(x).split(' ')[0] for x in s))),
        klient=('firma', lambda s: ', '.join(dict.fromkeys(str(x).split('/')[1].strip() for x in s if '/' in str(x)))), idx=('idx', 'first'), mesicni=('mesicni', 'last'))
    e['se_spz'] = e.index.map(si.groupby('spz').castka.sum()).fillna(0); e['vyf'] = e.se_spz + e.bez; e['chybi'] = e.ma - e.vyf
    e['model'] = e.idx.map(lambda i: re.sub(r'^rz:\s*\S+,\s*', '', str(V.loc[i, 'auto'])))
    mis = df[df.chybi > 0.5 * df.castka].groupby('spz').od.apply(lambda s: ', '.join(sorted({d.strftime('%m/%y') for d in s}, key=lambda x: (x[3:], x[:2]))))
    e['mesice'] = e.index.map(mis).fillna('')
    jj = J.groupby('spz').apply(lambda g: '; '.join(f'{r.odberatel} {fmt(r.castka)} Kč ({r.od}–{r.do})' for r in g.itertuples()), include_groups=False)
    e['jine'] = e.index.map(jj).fillna('')
    e['spol'] = (e.bez / e.vyf.where(e.vyf > 0, 1)).map(lambda x: 'nižší – část faktur bez SPZ, párováno podle šablony/částky' if x > 0.3 else 'dobrá – SPZ na fakturách')
    e.loc[e.vyf == 0, 'spol'] = 'žádná faktura na HAVEX'
    e.loc[(e.vyf == 0) & (e.do < pd.Timestamp('2021-09-15')), 'spol'] = 'nejistá – před 09/2021 se fakturovalo mimo Vyfakturuj'
    return e[e.chybi >= 1000].sort_values('chybi', ascending=False)

def sheet(ws, e, title, note):
    ws['A1'] = title; ws['A1'].font = Font(name=F, bold=True, size=14)
    ws['A2'] = note; ws['A2'].font = NOTE; ws['A2'].alignment = Alignment(wrap_text=True); ws.merge_cells('A2:M2'); ws.row_dimensions[2].height = 70
    cols = ['SPZ', 'Vůz', 'Klient (přehled)', 'Smlouvy', 'Od', 'Do', 'Měsíčně k fakturaci', 'Mělo být vyfakturováno', 'Vyfakturováno HAVEX', 'Chybí',
            'Měsíce bez faktury', 'Jiné faktury k autu (ne HAVEX)', 'Spolehlivost párování']
    for i, (c, w) in enumerate(zip(cols, [10, 32, 16, 22, 11, 11, 11, 13, 13, 12, 40, 44, 26]), 1):
        x = ws.cell(4, i, c); x.font = H; x.fill = HF; x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[4].height = 32
    r = 5
    for s, x in e.iterrows():
        vals = [s, x.model, x.klient, x.smlouvy, x.od.date(), x.do.date(), x.mesicni, round(x.ma), round(x.vyf), None, x.mesice, x.jine, x.spol]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v); c.font = T; c.border = THIN; c.alignment = Alignment(wrap_text=j in (11, 12), vertical='top')
            if j in (7, 8, 9): c.number_format = NUM
            if j in (5, 6): c.number_format = 'DD.MM.YYYY'
        c = ws.cell(r, 10, f'=H{r}-I{r}'); c.number_format = NUM; c.font = BO; c.border = THIN
        fill = YEL if x.jine else (RED if x.spol.startswith(('dobrá', 'žádná')) else None)
        if fill:
            for j in range(1, 14): ws.cell(r, j).fill = fill
        r += 1
    ws.cell(r, 1, 'Celkem').font = BO; ws.cell(r, 2, f'=COUNTA(A5:A{r-1})&" aut"').font = BO
    for j in (8, 9, 10):
        c = ws.cell(r, j, f'=SUM({get_column_letter(j)}5:{get_column_letter(j)}{r-1})'); c.font = BO; c.number_format = NUM
    ws.cell(r + 2, 1, 'Červeně = SPZ je na fakturách nebo na HAVEX nešla žádná faktura → rozdíl je spolehlivý. Žlutě = k autu existují faktury jiným odběratelům '
            '(prodej, přímý pronájem klientovi) → smlouva mohla skončit dřív nebo se fakturovalo jinak. Bez barvy = část faktur bez SPZ, rozdíl může být chybou párování.').font = NOTE
    ws.freeze_panes = 'B5'; ws.auto_filter.ref = f'A4:M{r-1}'

A = souhrn(P[P.od >= OD], SI[SI.duzp >= '2021-09'])
Z = souhrn(P[P.od < OD], SI[SI.duzp < '2021-09'])
wb = Workbook()
sheet(wb.active, A, 'Vozidla ke zpětné fakturaci HAVEX – nevyfakturováno (od 09/2021)',
      'Mělo být = sloupec „K fakturaci zpět na CC“ v přehledu zápůjček HAVEX (k 30. 9. 2026), po měsících od začátku do konce smlouvy, období začínající do 31. 8. 2026. '
      'Konec smlouvy je podle přehledu – pokud auto vrátili nebo prodali dřív, je „Mělo být“ nadsazené. Vyfakturováno = položky dokladů Vyfakturuj RX Auto CZ na HAVEX Mobility '
      '(bez pokut, jistin, škod a přefakturací): SPZ z textu položky, u starších faktur z hlavičky (text před položkami); zbylé položky bez SPZ přiřazené přes šablony pravidelných faktur (vč. ukončených) podle částky a dne začátku období. Rozdíly pod 1 000 Kč vynechány.')
wb.active.title = 'Od 09-2021'
sheet(wb.create_sheet('Před 09-2021'), Z, 'Zpětná fakturace HAVEX před 09/2021',
      'První faktura na HAVEX ve Vyfakturuj je z 15. 9. 2021. Do té doby se zpětná fakturace mohla vystavovat jinde (účetní, jiný systém) – ověřit v účetnictví, nemusí jít o nevyfakturované částky.')
out = B / 'output/Nevyfakturovano_HAVEX.xlsx'; wb.save(out)
print(out, len(A), round(A.chybi.sum()), '| pred', len(Z), round(Z.chybi.sum()))
for k, g in A.groupby('spol'): print(k, len(g), round(g.chybi.sum()))
print('jine', (A.jine != '').sum(), round(A[A.jine != ''].chybi.sum()))
print(A[A.spol.str.startswith(('dobrá', 'žádná'))].head(20)[['klient', 'od', 'do', 'ma', 'vyf', 'chybi', 'jine']].to_string())
