"""Přehled vozidel, která potřebují zimní pneu.

Vstupy (data/, mimo git – obsahují jména klientů):
  cz.json, sk.json      aktivní auta z CZ / SK Renteonu (api/car/getSimpleList, viz README)
  cz_inactive.json      neaktivní auta z CZ Renteonu (dohledání aut jen ve Vyfakturuj)
  vf_raw.json           detail aktivních pravidelných faktur z Vyfakturuj
  vf_next.json          {id pravidelné faktury: datum příští faktury}
Výstup: output/<datum>_Zimni_pneu_prehled.xlsx

Pravidla (Erik):
  - zimní pneu potřebuje auto s buybackem po 1.11.
  - Peugeot 3008 má celoroční pneu
  - vlastník podle štítku HVX / ACJ v InternalRemark
  - Renteon = krátkodobé, SPZ v pravidelné faktuře Vyfakturuj = dlouhodobý klient
  - ignorovat auta fakturovaná na HAVEX Mobility a FAM (Fábián)
  - start = opening card (první činnost vozu v Renteonu)
  - barevně auta, která končí v listopadu / prosinci daného roku
"""
import json, re, sys, datetime as dt
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).parent
DATA, OUT = BASE / 'data', BASE / 'output'
TODAY = dt.date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else dt.date.today()
YEAR = TODAY.year
CUT = dt.date(YEAR, 11, 1)
NOV, DEC = (dt.date(YEAR, 11, 1), dt.date(YEAR, 12, 1)), (dt.date(YEAR, 12, 1), dt.date(YEAR + 1, 1, 1))
FILL_NOV = PatternFill('solid', fgColor='F8CBAD')   # oranžová – konec v listopadu
FILL_DEC = PatternFill('solid', fgColor='FFE699')   # žlutá – konec v prosinci
F = 'Arial'

def pd(s):
    m = re.match(r'(\d{2})\.(\d{2})\.(\d{4})', s or '')
    return dt.date(int(m[3]), int(m[2]), int(m[1])) if m else None

def add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    y += d.year; m += 1
    import calendar
    return dt.date(y, m, min(d.day, calendar.monthrange(y, m)[1]))

def end_fill(d):
    if d and NOV[0] <= d < NOV[1]: return FILL_NOV
    if d and DEC[0] <= d < DEC[1]: return FILL_DEC
    return None

norm = lambda s: (s or '').replace(' ', '').upper()
load = lambda n: json.load(open(DATA / n, encoding='utf-8'))
cz, sk, ina = load('cz.json'), load('sk.json'), load('cz_inactive.json')
vf, nxt = load('vf_raw.json'), load('vf_next.json')

# --- Vyfakturuj: SPZ v pravidelných fakturách ---
PLATE = re.compile(r'\b(\d[A-Z][A-Z0-9]{4,5}|[A-Z]{2}\d{3}[A-Z]{2})\b')
SPACED = re.compile(r'\b(\d[A-Z]{2})\s(\d{4})\b')          # např. „8AR 9080“
for r in vf:
    txt = (r['name'] + ' ' + ' '.join(r['items'])).upper()
    txt = SPACED.sub(r'\1\2', txt)
    r['plates'] = sorted({x for x in PLATE.findall(txt) if re.search(r'\d{3}', x)})
    r['next'] = pd(nxt.get(r['id']))
    if 'počtu' in r['rtype']:          # vfpRepeatCount = zbývající počet dokladů
        cnt = int(r['cnt'] or 0)
        r['last'] = add_months(r['next'], cnt - 1) if r['next'] and cnt > 0 else None
    elif 'data' in r['rtype']:
        r['last'] = pd(r['end'])
    else:
        r['last'] = None               # bez ukončení
    r['open'] = r['last'] is None

EXCL = lambda c: 'HAVEX MOBILITY' in c.upper() or c.strip().upper() == 'FAM' or re.search('F[ÁA]BI[ÁA]N', c.upper())
EXCL_PLATES = {p for r in vf if EXCL(r['cust']) for p in r['plates']}
vf = [r for r in vf if not EXCL(r['cust'])]
vfmap = {}
for r in vf:
    for p in r['plates']: vfmap.setdefault(p, []).append(r)
fmt = lambda d: d.strftime('%d.%m.%Y') if d else 'bez ukončení'

def owner(rem):
    t = [x.strip() for x in (rem or '').upper().split('|')]
    return 'HAVEX' if 'HVX' in t else 'ACJ' if 'ACJ' in t else 'Jiný / vlastní'

RENT_ACT = lambda a: bool(a) and not re.search(r'vnitřní|interné|převod|líný|nečinn', a, re.I)

# --- Renteon CZ + SK ---
cars = {}
for c in cz: cars[norm(c['spz'])] = {'CZ': c}
for c in sk: cars.setdefault(norm(c['spz']), {})['SK'] = c

rows = []
for p, d in cars.items():
    if p in EXCL_PLATES: continue
    czc, skc = d.get('CZ'), d.get('SK')
    main = czc or skc
    flags = []
    if czc and skc:
        main = skc if (RENT_ACT(skc['act']) or not RENT_ACT(czc['act'])) and skc['act'] and not re.search('vnitřní', skc['act'], re.I) else czc
        if czc['bb'] != skc['bb']:
            flags.append(f"BB v CZ {czc['bb'] or '—'} vs SK {skc['bb'] or '—'} (použito pozdější)")
        if 'Peugeot' in czc['model'] and czc['model'].split()[1:2] != skc['model'].split()[1:2]:
            flags.append(f"Model CZ „{czc['model']}“ vs SK „{skc['model']}“")
    bbs = [pd(x['bb']) for x in (czc, skc) if x and pd(x['bb'])]
    firsts = [pd(x['first']) for x in (czc, skc) if x and pd(x.get('first'))]
    model = (czc or skc)['model']
    rem = (czc or skc)['rem'] or ''
    if skc and skc['rem'] and skc['rem'] != rem: rem += ' / SK: ' + skc['rem']
    vfr = vfmap.get(p, [])
    low = rem.lower()
    tire = [t for k, t in (('celoroč', 'Celoročky (poznámka v Renteonu)'), ('zimn', 'Má zimní (poznámka v Renteonu)'),
                           ('letní', 'Poznámka o letních pneu')) if k in low]
    rows.append(dict(
        spz=main['spz'], model=model, sipp=main['sipp'], office=main['office'],
        inst='+'.join(k for k in ('CZ', 'SK') if k in d), own=owner(rem),
        first=min(firsts) if firsts else None, bb=max(bbs) if bbs else None,
        status=main['status'], act=main['act'],
        typ='Dlouhodobý (Vyfakturuj)' if vfr else 'Krátkodobý (Renteon)',
        cust='; '.join(sorted({x['cust'].strip() for x in vfr})), vfname='; '.join(x['name'].strip() for x in vfr),
        vfend='; '.join(fmt(x['last']) for x in vfr), rem=rem, vin=main['vin'],
        allseason=bool(re.search(r'peugeot 3008', model, re.I)), tire='; '.join(tire), flags='; '.join(flags)))

need, allseas, early, other = [], [], [], []
for r in rows:
    if r['allseason']: allseas.append(r)
    elif r['own'] == 'Jiný / vlastní': other.append(r)
    elif r['bb'] and r['bb'] > CUT: need.append(r)
    else: early.append(r)
key = lambda r: (r['bb'] or dt.date.max, r['office'], r['model'], r['spz'])
hvx = sorted([r for r in need if r['own'] == 'HAVEX'], key=key)
acj = sorted([r for r in need if r['own'] == 'ACJ'], key=key)

# --- auta jen ve Vyfakturuj ---
inamap = {}
for c in ina: inamap.setdefault(norm(c['spz']), []).append(c)
vfonly = []
for p, lst in vfmap.items():
    if p in cars: continue
    item = next((re.sub(r'\(.*', '', it).strip() for x in lst for it in x['items'] if p in SPACED.sub(r'\1\2', it.upper())), '')
    ic = inamap.get(p)
    last = None if any(x['open'] for x in lst) else max(x['last'] for x in lst)
    starts = [pd(x['start']) for x in lst if pd(x['start'])]
    vfonly.append(dict(
        spz=p, item=item[:120], cust='; '.join(sorted({x['cust'].strip() for x in lst})),
        vfname='; '.join(x['name'].strip() for x in lst), start=min(starts) if starts else None,
        nxt=min((x['next'] for x in lst if x['next']), default=None), last=last,
        needs='ANO' if last is None or last > CUT else f'NE (fakturace končí do {CUT:%d.%m.})',
        p3008='ANO – celoroční' if '3008' in item else '',
        ren=(f"neaktivní v CZ Renteonu: {ic[0]['model']}, BB {ic[0]['bb']}, opening card {ic[0].get('first') or '?'}" if ic else 'není v Renteonu')))
vfonly.sort(key=lambda r: (r['last'] or dt.date.max, r['cust'], r['spz']))

# --- Excel ---
wb = Workbook()
HDR = PatternFill('solid', fgColor='1F3864')

def sheet(ws, headers, data, widths, title, end_col):
    ws['A1'] = title; ws['A1'].font = Font(name=F, bold=True, size=13)
    ws['A2'] = 'Oranžová = končí v listopadu, žlutá = končí v prosinci ' + str(YEAR)
    ws['A2'].font = Font(name=F, italic=True, size=9)
    ws.append(headers)
    for i in range(1, len(headers) + 1):
        c = ws.cell(3, i); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in data: ws.append(row)
    for row in ws.iter_rows(min_row=4):
        fill = end_fill(row[end_col - 1].value if isinstance(row[end_col - 1].value, dt.date) else None)
        for c in row:
            c.font = Font(name=F, size=10)
            if isinstance(c.value, dt.date): c.number_format = 'DD.MM.YYYY'
            if fill: c.fill = fill
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A4'
    if data:
        t = Table(displayName=re.sub(r'\W', '', ws.title) or 'T', ref=f"A3:{get_column_letter(len(headers))}{3 + len(data)}")
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
        ws.add_table(t)

H = ['SPZ', 'Model', 'SIPP', 'Pobočka', 'Renteon', 'Vlastník', 'Start (opening card)', 'Buyback', 'Typ pronájmu',
     'Odběratel (Vyfakturuj)', 'Pravidelná faktura', 'Poslední pravidelná fakt.', 'Stav v Renteonu', 'Poslední činnost',
     'Poznámka k pneu', 'Interní poznámka Renteon', 'Ke kontrole', 'VIN']
W = [11, 34, 7, 8, 8, 9, 12, 11, 22, 28, 30, 16, 12, 22, 26, 40, 40, 20]
BB_COL = 8
row = lambda r: [r['spz'], r['model'], r['sipp'], r['office'], r['inst'], r['own'], r['first'], r['bb'], r['typ'],
                 r['cust'], r['vfname'], r['vfend'], r['status'], r['act'], r['tire'], r['rem'], r['flags'], r['vin']]

ws = wb.active; ws.title = 'Přehled'
sheet(wb.create_sheet('HAVEX'), H, [row(r) for r in hvx], W, f'Zimní pneu – vozidla HAVEX (buyback po {CUT:%d.%m.%Y})', BB_COL)
sheet(wb.create_sheet('ACJ'), H, [row(r) for r in acj], W, f'Zimní pneu – vozidla ACJ (buyback po {CUT:%d.%m.%Y})', BB_COL)
HV = ['SPZ', 'Vozidlo (z položky faktury)', 'Odběratel', 'Pravidelná faktura', 'Start pravidelné fakt.', 'Příští faktura',
      'Poslední pravidelná fakt.', 'Potřeba zimních pneu', 'Peugeot 3008', 'Renteon']
sheet(wb.create_sheet('Mimo Renteon (Vyfakturuj)'), HV,
      [[r['spz'], r['item'], r['cust'], r['vfname'], r['start'], r['nxt'], r['last'] or 'bez ukončení', r['needs'], r['p3008'], r['ren']] for r in vfonly],
      [11, 50, 30, 34, 13, 13, 16, 22, 14, 60], 'Dlouhodobí klienti – vozidla jen ve Vyfakturuj (bez opening card, vlastník a buyback nejsou v Renteonu)', 7)
sheet(wb.create_sheet('Nepotřebují'), H + ['Důvod'],
      [row(r) + ['Peugeot 3008 – celoroční pneu'] for r in sorted(allseas, key=key)] +
      [row(r) + [f'Buyback do {CUT:%d.%m.%Y}' if r['bb'] else 'Bez buybacku'] for r in sorted(early, key=key)],
      W + [26], f'Nepotřebují zimní pneu – Peugeot 3008 (celoroční) a buyback do {CUT:%d.%m.%Y}', BB_COL)
sheet(wb.create_sheet('Jiný vlastník'), H, [row(r) for r in sorted(other, key=key)], W,
      'Vozidla bez štítku HVX/ACJ (vlastní / jiný vlastník) – rozhodnout ručně', BB_COL)

ws['A1'] = f'Zimní pneumatiky – přehled (stav k {TODAY:%d.%m.%Y})'; ws['A1'].font = Font(name=F, bold=True, size=14)
ws.append([])
ws.append(['', 'Krátkodobé (Renteon)', 'Dlouhodobé (Vyfakturuj)', 'Celkem', f'Končí v listopadu {YEAR}', f'Končí v prosinci {YEAR}'])
for c in ws[3]: c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HDR
rng = lambda sh, col, a, b: f'COUNTIFS({sh}!{col}:{col},">="&DATE({a.year},{a.month},1),{sh}!{col}:{col},"<"&DATE({b.year},{b.month},1))'
for i, name in enumerate(['HAVEX', 'ACJ'], start=4):
    ws.append([name, f'=COUNTIF({name}!I:I,"Krátkodobý*")', f'=COUNTIF({name}!I:I,"Dlouhodobý*")', f'=B{i}+C{i}',
               '=' + rng(name, 'H', *NOV), '=' + rng(name, 'H', *DEC)])
ws.append(['Celkem v Renteonu', '=B4+B5', '=C4+C5', '=D4+D5', '=E4+E5', '=F4+F5'])
MR = "'Mimo Renteon (Vyfakturuj)'"
ws.append(['Mimo Renteon – potřebují', '', f'=COUNTIF({MR}!H:H,"ANO")', '=C7', '=' + rng(MR, 'G', *NOV), '=' + rng(MR, 'G', *DEC)])
ws.append(['Celkem zimní pneu', '=B6', '=C6+C7', '=B8+C8', '=E6+E7', '=F6+F7'])
for r in ws.iter_rows(min_row=4, max_row=8):
    for c in r: c.font = Font(name=F, bold=(c.row == 8))
ws['E3'].fill = FILL_NOV; ws['F3'].fill = FILL_DEC
for c in (ws['E3'], ws['F3']): c.font = Font(name=F, bold=True)
notes = [
    'Pravidla:',
    '• Zdroj: aktivní vozidla v CZ Renteonu (rxauto.s3) a SK Renteonu (cwslovakia.s6), aktivní pravidelné faktury ve Vyfakturuj (firma RX Auto, s.r.o.).',
    f'• Zimní pneu potřebuje vozidlo s buybackem PO {CUT:%d.%m.%Y} (buyback přesně 1.11. se nepočítá).',
    '• Peugeot 3008 má celoroční pneu → list „Nepotřebují“.',
    '• Vlastník podle štítku HVX / ACJ v interní poznámce Renteonu.',
    '• Dlouhodobý = SPZ je v aktivní pravidelné faktuře ve Vyfakturuj; ostatní = krátkodobý pronájem z Renteonu.',
    '• Start (opening card) = datum první činnosti vozu v Renteonu; u aut v CZ i SK Renteonu je použito dřívější.',
    '• Barvy: oranžová = buyback (u aut mimo Renteon poslední pravidelná faktura) v listopadu, žlutá = v prosinci.',
    '• Vozidla v obou Renteonech jsou uvedena jednou; při rozdílném buybacku je použit pozdější a auto je označeno „Ke kontrole“.',
    '• Vynechána vozidla fakturovaná na HAVEX Mobility s.r.o. a FAM (Fábián).',
    '• Poslední pravidelná faktura: u ukončení „podle počtu dokladů“ = příští faktura + (zbývající počet − 1) měsíců.',
]
for i, n in enumerate(notes, start=10):
    ws.cell(i, 1, n).font = Font(name=F, bold=(i == 10), size=10)
ws.column_dimensions['A'].width = 28
for col in 'BCDEF': ws.column_dimensions[col].width = 22

OUT.mkdir(exist_ok=True)
out = OUT / f'{TODAY:%Y-%m-%d}_Zimni_pneu_prehled.xlsx'
wb.save(out)
print(out)
print('HAVEX', len(hvx), 'ACJ', len(acj), 'mimo Renteon', len(vfonly), sum(r['needs'] == 'ANO' for r in vfonly),
      'nepotřebují', len(allseas) + len(early), 'jiný vlastník', len(other))
for nm, lst, f in (('HAVEX', hvx, 'bb'), ('ACJ', acj, 'bb'), ('mimo', vfonly, 'last')):
    print(nm, 'listopad', sum(1 for r in lst if end_fill(r[f]) is FILL_NOV), 'prosinec', sum(1 for r in lst if end_fill(r[f]) is FILL_DEC))
print('bez opening card:', [r['spz'] for r in hvx + acj if not r['first']])
