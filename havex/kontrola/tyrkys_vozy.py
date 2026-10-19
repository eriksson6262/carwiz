"""Vytáhne tyrkysově podbarvená vozidla (sloupec E, FF00B0F0) z listu 'Vozy'
přehledu zápůjček HAVEX. A = start, B = konec, J = měsíční cena bez DPH (=G+H).

Použití: python3 tyrkys_vozy.py <prehled.xlsx> [od_konce=2026-08-01] > data/sheet.json
Výstup: JSON [[rz, start, konec, cena_mesic, radek, firma], ...]
"""
import json, re, sys
import openpyxl

TYRKYS = "FF00B0F0"

def main(path, od_konce="2026-08-01"):
    ws = openpyxl.load_workbook(path)["Vozy"]
    wv = openpyxl.load_workbook(path, data_only=True)["Vozy"]
    out = []
    for r in range(3, ws.max_row + 1):
        f = ws.cell(r, 5).fill
        if not (f and f.fill_type and f.fgColor.type == "rgb" and f.fgColor.rgb == TYRKYS):
            continue
        a, b, firma, auto, j = (wv.cell(r, c).value for c in (1, 2, 5, 6, 10))
        if not auto or not b:
            continue
        m = re.search(r"rz:\s*([A-Z0-9]+)", auto)
        if not m or str(b)[:10] < od_konce:
            continue
        out.append([m.group(1), str(a)[:10], str(b)[:10], j, r,
                    (firma or "").replace("RX Auto, s.r.o.", "").strip()])
    json.dump(out, sys.stdout, ensure_ascii=False)

if __name__ == "__main__":
    main(*sys.argv[1:])
