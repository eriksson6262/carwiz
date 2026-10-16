import json, sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
base=(sys.argv[2] if len(sys.argv)>2 else 'data/')
prot={r['rz']:r for r in json.load(open(base+'all.json'))}
vyf=json.load(open(base+'vyf.json')); ren=json.load(open(base+'renteon.json'))
con=json.load(open(base+'contracts.json'))
D=lambda s: dt.datetime.strptime(s,'%d.%m.%Y').date()
K1=('23.03.2026','22.09.2026'); K2=('31.03.2026','30.09.2026'); K3=('17.04.2026','16.10.2026')
rate={}
for rz in '8H84614 8H84604 8H84613 8H84624 8H84473 8H84627 8H84615'.split(): rate[rz]=(7585,)+K1
for rz in '8H84601 8H84617'.split(): rate[rz]=(7585,)+K2
for rz in '8H84621 8H84536 8H84540 8H84602 8H84623 8H84626 8H84594 8H84607'.split(): rate[rz]=(7193,)+K1
for rz in '8H84606 9H05923'.split(): rate[rz]=(7193,)+K2
rate['8H84694']=(11869,)+K2; rate['9H05927']=(12325,)+K2
rate['9H05926']=(12325,)+K3; rate['9H06025']=(8301,)+K3
for rz in '8H84595 8H84609 8H84593'.split(): rate[rz]=(8013,'24.03.2026','23.09.2026')
for rz in '8H84631 8H84666 8H84696 8H84669 8H84671'.split(): rate[rz]=(6537,'01.04.2026','30.09.2026')
rate['8H84674']=(6537,'10.04.2026','09.10.2026')
for rz in '8H84719 8H84726 8H84707'.split(): rate[rz]=(6537,'08.04.2026','07.10.2026')
rate['8H85167']=(8771,'29.09.2025','28.09.2026')
orig={'8H84473':'2600123','8H84593':'2600147','8H84609':'2600145','8H84595':'2600144'}
rows=[]
for rz,p in sorted(prot.items(), key=lambda x:(x[0] not in vyf, x[0])):
    r,od,do=rate[rz]; ret=D(p['vraceno'])
    months=12 if rz=='8H85167' else 6
    prol=max(0,(ret-D(do)).days) if rz in orig else 0
    fab=rz in vyf
    note=[]
    if rz in orig: note.append(f'prodloužení {con[rz]} od {D(do)+dt.timedelta(1):%d.%m.} do vrácení – sazba převzata z pův. smlouvy {orig[rz]} (předpoklad, ověřit dle PK)')
    if rz=='8H85167': note.append('splátkový kalendář 2500652 celkem 105 252 Kč bez DPH; vůz bez výnosu (soukromé užití)')
    if fab:
        n=vyf[rz]['rent']; o=vyf[rz]['damage']; f=vyf[rz]['fines']
        if prol: note.append('Fabianovi fakturováno jen 6 měsíců')
        src='Vyfakturuj'
    else:
        n=ren[rz]['rent']; o=round(ren[rz]['tot']-ren[rz]['rent'],2); f=None; src='Renteon'
    rows.append([rz,p['vuz'].strip(),'Fabian' if fab else 'RX Auto',con[rz],D(od),D(do),ret,r,months,prol,
                 None,p['havex_total'],None,n,o,None,None,None,f,src,'; '.join(note)])
wb=Workbook(); ws=wb.active; ws.title='Vozidla'
hdr=['RZ','Vůz','Provoz','Smlouva HAVEX','Začátek','Konec smlouvy','Vráceno','Měs. splátka HAVEX','Měsíců','Prodloužení (dny)',
     'Nájem HAVEX celkem','Škody/km dle protokolu HAVEX','NÁKLADY CELKEM','Výnos – nájem','Výnos – ostatní (škody / doplňky)','VÝNOSY CELKEM','ZISK','Marže','Pokuty přefakturované (info)','Zdroj výnosů','Poznámka']
F=Font(name='Arial',size=10); B=Font(name='Arial',size=10,bold=True); BL=Font(name='Arial',size=10,color='0000FF')
thin=Side(style='thin',color='BBBBBB'); bd=Border(top=thin,bottom=thin,left=thin,right=thin)
ws.append(['Náklady HAVEX vs. výnosy – vozidla z protokolů o vrácení (Beyr), stav k 16.10.2026, vše v Kč bez DPH']); ws['A1'].font=Font(name='Arial',size=12,bold=True)
ws.append([]); ws.append(hdr)
for c in range(1,len(hdr)+1):
    x=ws.cell(3,c); x.font=Font(name='Arial',size=10,bold=True,color='FFFFFF'); x.fill=PatternFill('solid',fgColor='1F3864'); x.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center'); x.border=bd
ws.row_dimensions[3].height=45
first=4
for i,row in enumerate(rows):
    R=first+i; ws.append(row)
    ws[f'K{R}']=f'=ROUND(H{R}*I{R}+H{R}/30*J{R},2)'
    ws[f'M{R}']=f'=K{R}+L{R}'; ws[f'P{R}']=f'=N{R}+O{R}'; ws[f'Q{R}']=f'=P{R}-M{R}'
    ws[f'R{R}']=f'=IF(P{R}=0,"–",Q{R}/P{R})'
    for c in range(1,len(hdr)+1):
        x=ws.cell(R,c); x.font=F; x.border=bd
    for c in 'HIJLNOS': ws[f'{c}{R}'].font=BL
    for c in 'EFG': ws[f'{c}{R}'].number_format='DD.MM.YYYY'
    for c in 'HKLMNOPQS': ws[f'{c}{R}'].number_format='#,##0;[Red]-#,##0'
    ws[f'R{R}'].number_format='0.0%;[Red]-0.0%'
    for c in 'MPQ': ws[f'{c}{R}'].font=B
last=first+len(rows)-1; T=last+1
ws[f'A{T}']='CELKEM'
for c in 'KLMNOPQS': ws[f'{c}{T}']=f'=SUM({c}{first}:{c}{last})'; ws[f'{c}{T}'].number_format='#,##0;[Red]-#,##0'
ws[f'R{T}']=f'=Q{T}/P{T}'; ws[f'R{T}'].number_format='0.0%;[Red]-0.0%'
for c in range(1,len(hdr)+1):
    x=ws.cell(T,c); x.font=B; x.fill=PatternFill('solid',fgColor='D9E1F2'); x.border=bd
from openpyxl.formatting.rule import CellIsRule
ws.conditional_formatting.add(f'Q{first}:Q{last}',CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor='F8CBAD')))
W=[10,34,9,12,11,11,11,11,7,10,13,13,14,13,14,14,13,8,12,11,70]
for i,w in enumerate(W,1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes='B4'; ws.auto_filter.ref=f'A3:U{last}'
# souhrn
s=wb.create_sheet('Souhrn')
s.append(['Souhrn podle provozu (Kč bez DPH)']); s['A1'].font=Font(name='Arial',size=12,bold=True); s.append([])
h2=['Provoz','Počet vozů','Nájem HAVEX','Škody HAVEX','Náklady celkem','Výnos nájem','Výnos ostatní','Výnosy celkem','Zisk','Marže','Ztrátové vozy']
s.append(h2)
for c in range(1,len(h2)+1):
    x=s.cell(3,c); x.font=Font(name='Arial',size=10,bold=True,color='FFFFFF'); x.fill=PatternFill('solid',fgColor='1F3864'); x.alignment=Alignment(wrap_text=True,horizontal='center')
rng=lambda c: f"Vozidla!${c}${first}:${c}${last}"
for j,g in enumerate(['Fabian','RX Auto'],4):
    s[f'A{j}']=g; s[f'B{j}']=f'=COUNTIF({rng("C")},A{j})'
    for col,src in zip('CDEFGHI','KLMNOPQ'): s[f'{col}{j}']=f'=SUMIF({rng("C")},$A{j},{rng(src)})'
    s[f'J{j}']=f'=IF(H{j}=0,0,I{j}/H{j})'; s[f'K{j}']=f'=COUNTIFS({rng("C")},A{j},{rng("Q")},"<0")'
s['A6']='CELKEM'
for col in 'BCDEFGHIK': s[f'{col}6']=f'=SUM({col}4:{col}5)'
s['J6']='=IF(H6=0,0,I6/H6)'
for R in range(4,7):
    for c in range(1,12):
        x=s.cell(R,c); x.font=B if R==6 else F; x.border=bd
        if c in (3,4,5,6,7,8,9): x.number_format='#,##0;[Red]-#,##0'
    s[f'J{R}'].number_format='0.0%;[Red]-0.0%'
notes=['Metodika:',
 '• Nájem HAVEX = měsíční splátka bez DPH × počet měsíců smlouvy (6, u 8H85167 12 dle platebního kalendáře) + u prodloužení splátka/30 × dny po konci smlouvy do vrácení.',
 '• Předčasné vrácení nesnižuje nájem (smlouva na pevnou dobu).',
 '• Škody/km HAVEX = celková částka z protokolu o vrácení (Beyr).',
 '• Výnosy Fabian = nájem fakturovaný ve Vyfakturuj (EUR × kurz faktury) + naše vyúčtování škod po vrácení; pokuty jsou průtokové, uvedeny jen pro info.',
 '• Výnosy RX Auto = Renteon, přehled příjmů dle vozu (smlouvy), nájem a celkem bez DPH.',
 '• Sazby prodloužení 2600687/689/691/692 převzaty z původních smluv – ověřit v podepsaných PK (DigiSign 15.10.2026).']
for i,t in enumerate(notes,9): s[f'A{i}']=t; s[f'A{i}'].font=B if i==9 else F
for i,w in enumerate([12,10,13,13,14,13,13,14,13,9,10],1): s.column_dimensions[get_column_letter(i)].width=w
wb.move_sheet('Souhrn',-1)
wb.save(sys.argv[1])
